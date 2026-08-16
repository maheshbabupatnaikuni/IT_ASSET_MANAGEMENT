import io
from collections import Counter
from datetime import date, datetime, time, timedelta, timezone
from html import escape
from zoneinfo import ZoneInfo
from pathlib import Path
import pandas as pd
from flask import Blueprint, render_template, send_file, request, abort, redirect, url_for, current_app, session
from sqlalchemy import and_, or_
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from ..extensions import db
from ..models import Asset, MovementLog, RepairLog, Employee, Department, Location, Manufacturer, Vendor, AssetCategory, AssetAssignment, AssetLifecycle, DynamicFieldValue, ImportSourceLink, ImportBatch, InfrastructureItem, AssetRequest, ApprovalRequest, AuditLog, User
from ..services.audit import audit
from ..services.reporting import asset_row, employee_row, infrastructure_row, display_value, model_columns, rectangular_rows
from ..services.lbac import workspace_location_id
from .helpers import login_required, normalized_role, can
bp=Blueprint("reports",__name__,url_prefix="/reports")
LOCAL_TZ=ZoneInfo("Asia/Kolkata")
def branded_pdf_heading(title, styles):
 elements=[
  Paragraph("<b><font color='#203F78'>RAY</font> <font color='#1479B8'>IT</font> <font color='#203F78'>MANAGEMENT</font></b>",styles["Heading2"]),
  Spacer(1,8),Paragraph(title,styles["Title"]),Spacer(1,12),
 ]
 return elements
@bp.before_request
def restrict_reports():
 if not normalized_role():return redirect(url_for("auth.login"))
 if request.endpoint=="reports.search":
  if not any(can(code) for code in ("asset.view","employee.view","infrastructure.view")):abort(403)
  return
 if request.endpoint in ("reports.archive_preview","reports.archive_pdf"):
  required="database.backup"
 elif request.endpoint=="reports.inventory_print_pdf":
  required="report.print"
 else:
  required="report.export" if request.endpoint in ("reports.inventory","reports.dataset") else ("assignment.acknowledgement.download" if request.endpoint=="reports.acknowledgement" else "report.view")
 if not can(required):abort(403)

def report_date_window(required=False):
 try:
  start_date=date.fromisoformat(request.args.get("date_from","")) if request.args.get("date_from") else None
  end_date=date.fromisoformat(request.args.get("date_to","")) if request.args.get("date_to") else None
 except ValueError:
  abort(400,description="Enter valid From and To dates.")
 if required and (not start_date or not end_date):abort(400,description="Both From Date and To Date are required.")
 if start_date and end_date and start_date>end_date:abort(400,description="From Date cannot be later than To Date.")
 start_utc=datetime.combine(start_date,time.min,LOCAL_TZ).astimezone(timezone.utc).replace(tzinfo=None) if start_date else None
 end_utc=datetime.combine(end_date+timedelta(days=1),time.min,LOCAL_TZ).astimezone(timezone.utc).replace(tzinfo=None) if end_date else None
 return start_date,end_date,start_utc,end_utc

def apply_changed_window(query,model):
 _from,_to,start_utc,end_utc=report_date_window()
 if start_utc or end_utc:
  windows=[]
  for column in (model.created_at,model.updated_at):
   bounds=[]
   if start_utc:bounds.append(column>=start_utc)
   if end_utc:bounds.append(column<end_utc)
   windows.append(and_(*bounds))
  query=query.filter(or_(*windows))
 return query

def apply_event_window(query,column):
 _from,_to,start_utc,end_utc=report_date_window()
 if start_utc:query=query.filter(column>=start_utc)
 if end_utc:query=query.filter(column<end_utc)
 return query
def filtered_assets():
 query=Asset.query
 for key,column in [("status",Asset.status),("category",Asset.category_id),("department",Asset.department_id),("location",Asset.location_id),("vendor",Asset.vendor_id),("employee",Asset.employee_id)]:
  if request.args.get(key):query=query.filter(column==request.args[key])
 query=apply_changed_window(query,Asset)
 if request.args.get("source"):query=query.join(ImportSourceLink,ImportSourceLink.asset_id==Asset.id).filter(ImportSourceLink.batch_id==request.args["source"],ImportSourceLink.detached==False)
 if request.args.get("sheet"):query=query.join(ImportSourceLink,ImportSourceLink.asset_id==Asset.id).filter(ImportSourceLink.sheet_name==request.args["sheet"],ImportSourceLink.detached==False)
 return query.distinct().order_by(Asset.asset_tag)
def inventory_rows(query=None):
 assets=(query or Asset.query.order_by(Asset.asset_tag)).all()
 return [asset_row(asset) for asset in assets]

def inventory_summary_data(assets):
 location_id=request.args.get("location",type=int) or workspace_location_id()
 location=db.session.get(Location,location_id) if location_id else None
 employee_query=Employee.query
 infrastructure_query=InfrastructureItem.query
 request_query=AssetRequest.query
 if location_id:
  employee_query=employee_query.filter(Employee.location_id==location_id)
  infrastructure_query=infrastructure_query.filter(InfrastructureItem.location_id==location_id)
  request_query=request_query.join(Employee,AssetRequest.requested_for_id==Employee.id).filter(Employee.location_id==location_id)
 employees=employee_query.all()
 status_counts=Counter((asset.status or "Not Specified").strip() for asset in assets)
 active_employees=sum(1 for employee in employees if (employee.employment_status or "Active").strip().lower()=="active")
 employee_ids={asset.employee_id for asset in assets if asset.employee_id}
 pending_approvals=request_query.filter(AssetRequest.status.in_(("HR Raised","Pending Approval","Admin Processed","Returned to Admin"))).count()
 return {
  "location_name":location.name if location else "All Locations",
  "status_counts":status_counts,
  "active_employees":active_employees,
  "employees_with_assets":len(employee_ids),
  "infrastructure":infrastructure_query.count(),
  "pending_approvals":pending_approvals,
  "categories":Counter(asset.category.name if asset.category else "Not Specified" for asset in assets),
  "departments":Counter(asset.employee.department.name if asset.employee and asset.employee.department else (asset.department.name if asset.department else "Not Specified") for asset in assets),
 }

def employee_rows():
 query=Employee.query
 for key,column in (("department",Employee.department_id),("location",Employee.location_id),("status",Employee.employment_status)):
  if request.args.get(key):query=query.filter(column==request.args[key])
 query=apply_changed_window(query,Employee)
 return [employee_row(employee) for employee in query.order_by(Employee.name).all()]

def infrastructure_rows():
 query=InfrastructureItem.query
 for key,column in (("location",InfrastructureItem.location_id),("department",InfrastructureItem.department_id),("vendor",InfrastructureItem.vendor_id),("status",InfrastructureItem.status)):
  if request.args.get(key):query=query.filter(column==request.args[key])
 query=apply_changed_window(query,InfrastructureItem)
 return [infrastructure_row(item) for item in query.order_by(InfrastructureItem.infrastructure_reference).all()]
@bp.route("/")
@login_required
def index(): return render_template("reports_v2.html",assets=Asset.query.count(),infrastructure=InfrastructureItem.query.count(),movements=MovementLog.query.count(),repairs=RepairLog.query.count(),categories=AssetCategory.query.all(),departments=Department.query.all(),locations=Location.query.filter_by(active=True).order_by(Location.name).all(),vendors=Vendor.query.all(),employees=Employee.query.all(),sources=ImportBatch.query.order_by(ImportBatch.filename).all())
@bp.route("/inventory/<fmt>")
@login_required
def inventory(fmt):
 if fmt!="download": return redirect(url_for("reports.preview",kind="inventory",**request.args))
 rows=inventory_rows(filtered_assets()); df=pd.DataFrame(rows)
 return send_rows(rows,"excel","IT Asset Inventory","inventory")

def dataset_rows(kind):
 if kind=="inventory": return inventory_rows(filtered_assets())
 if kind=="employees":return employee_rows()
 if kind=="repairs":
  query=apply_changed_window(RepairLog.query,RepairLog)
  return [{**model_columns(row,exclude={"id","asset_id","vendor_id"}),"Asset Reference":row.asset.display_reference,"Asset Tag":row.asset.display_asset_tag or row.asset.asset_tag,"Asset Status":row.asset.status,"Location":row.asset.location.name if row.asset.location else "","Vendor":row.vendor.name if row.vendor else ""} for row in query.order_by(RepairLog.repair_date.desc()).all()]
 if kind=="assignments":
  query=apply_event_window(AssetAssignment.query,AssetAssignment.assigned_at)
  return [{**model_columns(row,exclude={"id","asset_id","employee_id"}),"Asset Reference":row.asset.display_reference,"Asset Tag":row.asset.display_asset_tag or row.asset.asset_tag,"Asset Category":row.asset.category.name if row.asset.category else "","Employee Code":row.employee.employee_code,"Employee Name":row.employee.name,"Department":row.employee.department.name if row.employee.department else "","Location":row.employee.location.name if row.employee.location else ""} for row in query.order_by(AssetAssignment.assigned_at.desc()).all()]
 if kind=="history":
  query=apply_event_window(AssetLifecycle.query,AssetLifecycle.occurred_at)
  return [{**model_columns(row,exclude={"id","asset_id","from_employee_id","to_employee_id","location_id"}),"Asset Reference":row.asset.display_reference,"Asset Tag":row.asset.display_asset_tag or row.asset.asset_tag,"From Employee":row.from_employee.name if row.from_employee else "","To Employee":row.to_employee.name if row.to_employee else "","Location":row.location.name if row.location else ""} for row in query.order_by(AssetLifecycle.occurred_at.desc()).all()]
 if kind=="movements":
  query=apply_event_window(MovementLog.query,MovementLog.date)
  return [{**model_columns(row,exclude={"id","asset_id"}),"Asset Reference":row.asset.display_reference,"Asset Tag":row.asset.display_asset_tag or row.asset.asset_tag} for row in query.order_by(MovementLog.date.desc()).all()]
 if kind=="infrastructure":return infrastructure_rows()
 abort(404)

@bp.route("/preview/<kind>")
@login_required
def preview(kind):
 rows=dataset_rows(kind)
 columns,rows=rectangular_rows(rows)
 titles={"inventory":"Complete IT Asset Inventory","employees":"Complete Employee Register","infrastructure":"Complete Infrastructure Register","assignments":"Assignment History","repairs":"Repair History","history":"Asset Lifecycle History","movements":"Movement History"}
 return render_template("report_preview.html",title=titles.get(kind,kind.title()),kind=kind,rows=rows,columns=columns,generated_at=datetime.now(LOCAL_TZ))

@bp.route("/search")
@login_required
def search():
 q=request.args.get("q","").strip()
 assets=[]; employees=[]; infrastructure=[]
 if can("asset.view"):
  query=Asset.query.outerjoin(Asset.employee).outerjoin(Asset.department).outerjoin(Asset.location).outerjoin(Asset.manufacturer).outerjoin(Asset.vendor)
  if q:
   dynamic_ids=db.session.query(DynamicFieldValue.asset_id).filter(DynamicFieldValue.value.contains(q))
   source_ids=db.session.query(ImportSourceLink.asset_id).join(ImportBatch).filter(or_(ImportBatch.filename.contains(q),ImportSourceLink.sheet_name.contains(q)))
   query=query.filter(or_(Asset.location_asset_reference.contains(q),Asset.system_asset_reference.contains(q),Asset.asset_tag.contains(q),Asset.source_asset_tag.contains(q),Asset.serial_number.contains(q),Asset.source_serial_number.contains(q),Asset.model.contains(q),Asset.invoice_number.contains(q),Asset.ip_address.contains(q),Asset.mac_address.contains(q),Asset.hostname.contains(q),Asset.status.contains(q),Employee.name.contains(q),Employee.employee_code.contains(q),Department.name.contains(q),Location.name.contains(q),Manufacturer.name.contains(q),Vendor.name.contains(q),Asset.id.in_(dynamic_ids),Asset.id.in_(source_ids)))
  for key,column in [("status",Asset.status),("category",Asset.category_id),("department",Asset.department_id),("condition",Asset.condition),("location",Asset.location_id),("vendor",Asset.vendor_id)]:
   if request.args.get(key):query=query.filter(column==request.args[key])
  assets=query.order_by(Asset.asset_tag).limit(100).all()
 if can("employee.view") and q:
  employees=Employee.query.join(Employee.department).join(Employee.location).filter(or_(Employee.name.contains(q),Employee.employee_code.contains(q),Employee.email.contains(q),Employee.designation.contains(q),Department.name.contains(q),Location.name.contains(q))).order_by(Employee.name).limit(50).all()
 if can("infrastructure.view") and q:
  infrastructure=InfrastructureItem.query.outerjoin(InfrastructureItem.manufacturer).outerjoin(InfrastructureItem.vendor).outerjoin(InfrastructureItem.department).filter(or_(InfrastructureItem.infrastructure_reference.contains(q),InfrastructureItem.name.contains(q),InfrastructureItem.asset_tag.contains(q),InfrastructureItem.serial_number.contains(q),InfrastructureItem.model.contains(q),InfrastructureItem.management_ip.contains(q),InfrastructureItem.ip_address.contains(q),InfrastructureItem.mac_address.contains(q),InfrastructureItem.site.contains(q),InfrastructureItem.building.contains(q),InfrastructureItem.floor.contains(q),InfrastructureItem.room.contains(q),Manufacturer.name.contains(q),Vendor.name.contains(q),Department.name.contains(q))).order_by(InfrastructureItem.infrastructure_reference).limit(50).all()
 return render_template("search.html",assets=assets,employees=employees,infrastructure=infrastructure,q=q,categories=AssetCategory.query.all(),departments=Department.query.all(),locations=Location.query.all(),vendors=Vendor.query.all())

@bp.route("/acknowledgement/<kind>/<int:item_id>")
@login_required
def acknowledgement(kind,item_id):
 asset=None;employee=None
 if kind=="clearance":employee=Employee.query.get_or_404(item_id);assets=Asset.query.filter_by(employee_id=employee.id).all()
 else:
  asset=Asset.query.get_or_404(item_id);latest=AssetAssignment.query.filter_by(asset_id=asset.id).order_by(AssetAssignment.assigned_at.desc()).first();employee=asset.employee or (latest.employee if latest else None);assets=[asset]
 out=io.BytesIO();doc=SimpleDocTemplate(out,pagesize=A4);styles=getSampleStyleSheet();title={"assignment":"Asset Assignment Acknowledgement","return":"Asset Return Acknowledgement","transfer":"Asset Transfer Acknowledgement","clearance":"Employee IT Clearance"}.get(kind,"Asset Acknowledgement");body=branded_pdf_heading(title,styles)+[Paragraph(f"Employee: {employee.name if employee else 'Unassigned'}",styles["Normal"]),Spacer(1,12)];data=[["Asset Tag","Category","Serial","Condition","Accessories"]]+[[(a.source_asset_tag or a.asset_tag),a.category.name,(a.source_serial_number or a.serial_number or ""),a.condition or "","See assignment record"] for a in assets];table=Table(data,repeatRows=1);table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#203f78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("LINEBELOW",(0,0),(-1,0),2,colors.HexColor("#76b82a"))]));body += [table,Spacer(1,48),Paragraph("Employee Signature: ____________________",styles["Normal"]),Spacer(1,28),Paragraph("IT Representative: ____________________",styles["Normal"]),Spacer(1,28),Paragraph("Date: ____________________",styles["Normal"])];doc.build(body);out.seek(0);return send_file(out,as_attachment=False,download_name=f"{kind}_acknowledgement.pdf",mimetype="application/pdf")

def archive_sections():
 start_date,end_date,_start_utc,_end_utc=report_date_window(required=True)
 assets=[asset_row(row) for row in apply_changed_window(Asset.query,Asset).order_by(Asset.asset_tag).all()]
 employees=[employee_row(row) for row in apply_changed_window(Employee.query,Employee).order_by(Employee.name).all()]
 infrastructure=[infrastructure_row(row) for row in apply_changed_window(InfrastructureItem.query,InfrastructureItem).order_by(InfrastructureItem.infrastructure_reference).all()]
 assignments=dataset_rows("assignments")
 repairs=dataset_rows("repairs")
 lifecycle=dataset_rows("history")
 movements=dataset_rows("movements")
 assignment_requests=[]
 for row in apply_changed_window(AssetRequest.query,AssetRequest).order_by(AssetRequest.created_at).all():
  item=model_columns(row,exclude={"requested_for_id","category_id","requested_by_id","asset_id","processed_by_id","approved_by_id"})
  item.update({"Employee":row.requested_for.name if row.requested_for else "","Category":row.category.name if row.category else "","Asset":row.asset.display_reference if row.asset else "","Requested By":row.requested_by.full_name if row.requested_by else "","Processed By":row.processed_by.full_name if row.processed_by else "","Approved By":row.approved_by.full_name if row.approved_by else ""})
  assignment_requests.append(item)
 approvals=[]
 for row in apply_changed_window(ApprovalRequest.query,ApprovalRequest).order_by(ApprovalRequest.created_at).all():
  item=model_columns(row,exclude={"location_id","target_location_id"});item["Location"]=row.location.name if row.location else "";item["Target Location"]=row.target_location.name if row.target_location else "";approvals.append(item)
 audit_rows=[]
 query=apply_event_window(AuditLog.query,AuditLog.timestamp)
 for row in query.order_by(AuditLog.timestamp).all():
  audit_rows.append({"Date and Time":display_value(row.timestamp),"User":row.user,"Entity":row.entity,"Entity ID":row.entity_id,"Action":row.action,"Location":row.location.name if row.location else "","Reason":row.reason,"Old Values":row.old_value,"New Values":row.new_value})
 users=[]
 for row in apply_changed_window(User.query,User).order_by(User.username).all():
  users.append({"Username":row.username,"Full Name":row.full_name,"Role":row.role.name if row.role else "","Status":"Active" if row.active else "Inactive","Approval Level":row.approval_level,"Assigned Locations":"; ".join(location.name for location in row.locations),"Created":display_value(row.created_at),"Updated":display_value(row.updated_at),"Last Login":display_value(row.last_login_at)})
 master_rows=[]
 for model,name in ((Location,"Location"),(Department,"Department"),(AssetCategory,"Asset Category"),(Manufacturer,"Manufacturer"),(Vendor,"Vendor")):
  for row in apply_changed_window(model.query,model).order_by(model.name).all():
   item=model_columns(row,exclude={"id"});item["Master Type"]=name;master_rows.append(item)
 return start_date,end_date,[
  ("Assets",assets),("Employees",employees),("Infrastructure",infrastructure),
  ("Assignments",assignments),("Repairs",repairs),("Asset Lifecycle",lifecycle),
  ("Movement History",movements),("Assignment Requests",assignment_requests),
  ("Action and Transfer Approvals",approvals),("User Directory - No Password Data",users),
  ("Master Data",master_rows),("Audit Trail",audit_rows),
 ]

def archive_page(canvas,doc):
 canvas.saveState();width,height=A4
 canvas.setFont("Helvetica-Bold",9);canvas.setFillColor(colors.HexColor("#203f78"));canvas.drawString(doc.leftMargin,height-28,"IT Asset Management")
 canvas.setStrokeColor(colors.HexColor("#76b82a"));canvas.line(doc.leftMargin,height-34,width-doc.rightMargin,height-34)
 canvas.setFont("Helvetica",8);canvas.setFillColor(colors.HexColor("#5f6b7a"));canvas.drawRightString(width-doc.rightMargin,20,f"Page {doc.page}")
 canvas.restoreState()

def inventory_pdf_page(canvas,doc):
 canvas.saveState();width,height=A4
 canvas.setFont("Helvetica",7);canvas.setFillColor(colors.HexColor("#6b7280"))
 canvas.drawRightString(width-doc.rightMargin,19,f"Page {doc.page}")
 canvas.restoreState()

def inventory_document_header(location_name,generated_label,styles):
 header=Table([[
  Paragraph("<b><font color='#203f78'>IT Asset Management</font></b>",styles["InventoryHeader"]),
  Paragraph(f"<para align='right'><b>{escape(location_name)} Asset Inventory</b><br/><font size='6.8'>Generated {escape(generated_label)}</font></para>",styles["InventoryHeader"]),
 ]],colWidths=[260,263])
 header.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LINEBELOW",(0,0),(-1,-1),1.2,colors.HexColor("#76b82a")),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),0),("TOPPADDING",(0,0),(-1,-1),0),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
 return [header,Spacer(1,18)]

def inventory_summary_flowables(assets,summary,styles,generated_label):
 status_counts=summary["status_counts"]
 def status_count(*names):
  wanted={name.lower() for name in names}
  return sum(value for name,value in status_counts.items() if name.lower() in wanted)
 kpis=[
  ("Total Assets",len(assets)),("Available",status_count("Available")),("Assigned",status_count("Assigned")),
  ("Under Repair",status_count("Under Repair","Repair")),("Blocked",status_count("Blocked")),("Lost",status_count("Lost")),
  ("Scrapped",status_count("Scrapped","Disposed")),("Active Employees",summary["active_employees"]),("Employees Holding Assets",summary["employees_with_assets"]),
  ("Pending Approvals",summary["pending_approvals"]),("Infrastructure",summary["infrastructure"]),("Other Asset States",len(assets)-sum(status_count(name) for name in ("Available","Assigned","Under Repair","Repair","Blocked","Lost","Scrapped","Disposed"))),
 ]
 card_rows=[]
 for offset in range(0,len(kpis),4):
  card_rows.append([Paragraph(f"<font size='16'><b>{value}</b></font><br/><font size='7.5'>{escape(label)}</font>",styles["InventoryCard"]) for label,value in kpis[offset:offset+4]])
 cards=Table(card_rows,colWidths=[128,128,128,128],rowHeights=48)
 cards.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#eef4f8")),("BOX",(0,0),(-1,-1),0.45,colors.HexColor("#9fb4c7")),("INNERGRID",(0,0),(-1,-1),3.5,colors.white),("VALIGN",(0,0),(-1,-1),"MIDDLE")]))
 def breakdown(title,counter):
  entries=compact_breakdown(counter,8)
  data=[[Paragraph(f"<b><font color='#ffffff'>{escape(title)}</font></b>",styles["InventoryCell"]),Paragraph("<font color='#ffffff'>Count</font>",styles["InventoryCell"])]]
  data.extend([Paragraph(escape(str(name)),styles["InventoryCell"]),str(value)] for name,value in entries)
  table=Table(data,colWidths=[202,45],repeatRows=1)
  table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#203f78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#cbd5e1")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f5f8fb")]),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3)]))
  return table
 breakdowns=Table([[breakdown("Assets by Category",summary["categories"]),breakdown("Assigned Assets by Employee Department",Counter(asset.employee.department.name if asset.employee and asset.employee.department else "Not Assigned" for asset in assets if (asset.status or "").lower()=="assigned"))]],colWidths=[256,256])
 breakdowns.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),8)]))
 return inventory_document_header(summary["location_name"],generated_label,styles)+[Paragraph(f"{escape(summary['location_name'])} - Asset Inventory Summary",styles["InventoryTitle"]),Paragraph(f"Dashboard-style summary for the exact filtered inventory. Detailed asset profiles begin on page 2.",styles["InventorySubTitle"]),Spacer(1,10),cards,Spacer(1,12),breakdowns]

def asset_profile_flowables(asset,index,total,styles,location_name,generated_label):
 row=asset_row(asset)
 populated=[(str(key),display_value(value)) for key,value in row.items() if value not in (None,"")]
 # Four columns normally; use three field/value pairs for unusually field-rich records.
 pairs_per_row=3 if len(populated)>48 else 2
 font_size=5.7 if pairs_per_row==3 else (6.2 if len(populated)>38 else 6.8)
 leading=font_size+1.4
 cell_style=ParagraphStyle(f"AssetCell{index}",parent=styles["InventoryCell"],fontSize=font_size,leading=leading)
 rows=[]
 for offset in range(0,len(populated),pairs_per_row):
  values=[]
  for key,value in populated[offset:offset+pairs_per_row]:
   values.extend((Paragraph(f"<b>{escape(key)}</b>",cell_style),Paragraph(escape(str(value)).replace("\n","<br/>"),cell_style)))
  while len(values)<pairs_per_row*2:values.extend(("",""))
  rows.append(values)
 usable_width=523
 label_width=75 if pairs_per_row==3 else 92
 value_width=(usable_width-(label_width*pairs_per_row))/pairs_per_row
 table=Table(rows,colWidths=sum(([label_width,value_width] for _ in range(pairs_per_row)),[]),splitByRow=0)
 table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("BACKGROUND",(0,0),(-1,-1),colors.white),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#eef4f8")),("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#cbd5e1")),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2)] + [("BACKGROUND",(column,0),(column,-1),colors.HexColor("#eef4f8")) for column in range(0,pairs_per_row*2,2)]))
 title=asset.display_reference or asset.display_asset_tag or asset.asset_tag
 subtitle=" | ".join(value for value in (asset.category.name if asset.category else None,asset.status,asset.location.name if asset.location else None) if value)
 # A leading spacer keeps the repeated document header inside the printable
 # frame after an explicit PageBreak in all ReportLab/Windows builds.
 return [Spacer(1,22)]+inventory_document_header(location_name,generated_label,styles)+[Paragraph(f"Asset {index} of {total}: {escape(title)}",styles["AssetTitle"]),Paragraph(escape(subtitle),styles["InventorySubTitle"]),Spacer(1,7),table]

def build_inventory_print_pdf(assets):
 output=io.BytesIO();styles=getSampleStyleSheet()
 styles.add(ParagraphStyle(name="InventoryTitle",parent=styles["Title"],fontSize=20,leading=23,textColor=colors.HexColor("#203f78"),spaceAfter=4))
 styles.add(ParagraphStyle(name="AssetTitle",parent=styles["Heading1"],fontSize=14,leading=17,textColor=colors.HexColor("#203f78"),spaceAfter=3))
 styles.add(ParagraphStyle(name="InventorySubTitle",parent=styles["BodyText"],fontSize=8,leading=10,textColor=colors.HexColor("#526071")))
 styles.add(ParagraphStyle(name="InventoryCell",parent=styles["BodyText"],fontSize=7,leading=8.5))
 styles.add(ParagraphStyle(name="InventoryCard",parent=styles["BodyText"],alignment=TA_CENTER,textColor=colors.HexColor("#203f78"),leading=10))
 styles.add(ParagraphStyle(name="InventoryHeader",parent=styles["BodyText"],fontSize=9,leading=10,textColor=colors.HexColor("#526071")))
 summary=inventory_summary_data(assets)
 doc=SimpleDocTemplate(output,pagesize=A4,rightMargin=36,leftMargin=36,topMargin=30,bottomMargin=34,title=f"{summary['location_name']} Asset Inventory")
 generated_label=datetime.now(LOCAL_TZ).strftime("%d %b %Y %I:%M:%S %p")
 story=inventory_summary_flowables(assets,summary,styles,generated_label)
 for index,asset in enumerate(assets,1):
  story.append(PageBreak());story.extend(asset_profile_flowables(asset,index,len(assets),styles,summary["location_name"],generated_label))
 if not assets:story.extend([Spacer(1,18),Paragraph("No assets match the selected filters.",styles["InventorySubTitle"])])
 doc.build(story,onFirstPage=inventory_pdf_page,onLaterPages=inventory_pdf_page);output.seek(0)
 return output,summary

@bp.route("/inventory/print-preview.pdf")
@login_required
def inventory_print_pdf():
 assets=filtered_assets().all();output,summary=build_inventory_print_pdf(assets)
 filename=f"{summary['location_name'].replace(' ','_')}_Asset_Inventory_{datetime.now(LOCAL_TZ):%Y%m%d}.pdf"
 return send_file(output,as_attachment=False,download_name=filename,mimetype="application/pdf")

def summary_frequency(rows,key):
 return Counter(str(row.get(key) or "Not Specified") for row in rows)

def compact_breakdown(counter,limit=7):
 rows=counter.most_common(limit)
 remaining=sum(counter.values())-sum(value for _name,value in rows)
 if remaining:rows.append(("Other",remaining))
 return rows

def archive_summary_flowables(sections,styles):
 by_name={title:rows for title,rows in sections}
 assets=by_name.get("Assets",[]);employees=by_name.get("Employees",[]);infrastructure=by_name.get("Infrastructure",[])
 assignments=by_name.get("Assignments",[]);repairs=by_name.get("Repairs",[]);requests_rows=by_name.get("Assignment Requests",[]);approvals=by_name.get("Action and Transfer Approvals",[])
 status_counts=summary_frequency(assets,"Status")
 pending_approvals=sum(1 for row in requests_rows+approvals if str(row.get("Status") or "").lower().startswith("pending"))
 active_employees=sum(1 for row in employees if str(row.get("Employment Status") or "").lower()=="active")
 open_repairs=sum(1 for row in repairs if str(row.get("Status") or "").lower() in ("sent","open","under repair"))
 kpis=[
  ("Assets",len(assets)),("Assigned",status_counts.get("Assigned",0)),("Available",status_counts.get("Available",0)),
  ("Active Employees",active_employees),("Infrastructure",len(infrastructure)),("Assignments",len(assignments)),
  ("Repairs",len(repairs)),("Open Repairs",open_repairs),("Pending Approvals",pending_approvals),
 ]
 card_rows=[]
 for offset in range(0,len(kpis),3):
  card_rows.append([Paragraph(f"<font size='18'><b>{value}</b></font><br/><font size='8'>{escape(label)}</font>",styles["SummaryCard"]) for label,value in kpis[offset:offset+3]])
 cards=Table(card_rows,colWidths=[153,153,153],rowHeights=48)
 cards.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#eef4f8")),("BOX",(0,0),(-1,-1),0.5,colors.HexColor("#9fb4c7")),("INNERGRID",(0,0),(-1,-1),4,colors.white),("VALIGN",(0,0),(-1,-1),"MIDDLE")]))

 def breakdown_table(title,rows,width=220):
  data=[[Paragraph(f"<b>{escape(title)}</b>",styles["ArchiveCell"]),Paragraph("Count",styles["ArchiveCell"])]]
  data += [[Paragraph(escape(str(name)),styles["ArchiveCell"]),str(value)] for name,value in rows]
  table=Table(data,colWidths=[width-45,45],repeatRows=1)
  table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#203f78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#cbd5e1")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f5f8fb")])]))
  return table

 operations=[("Assignments",len(assignments)),("Repairs",len(repairs)),("Lifecycle Events",len(by_name.get("Asset Lifecycle",[]))),("Movements",len(by_name.get("Movement History",[]))),("Assignment Requests",len(requests_rows)),("Other Approvals",len(approvals)),("Audit Entries",len(by_name.get("Audit Trail",[])))]
 first_pair=Table([[breakdown_table("Asset Status",compact_breakdown(status_counts)),breakdown_table("Operations in Period",operations)]],colWidths=[230,230])
 first_pair.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),8)]))

 location_assets=summary_frequency(assets,"Location");location_employees=summary_frequency(employees,"Location");location_infra=summary_frequency(infrastructure,"Company Location")
 locations=sorted(set(location_assets)|set(location_employees)|set(location_infra),key=lambda name:-(location_assets[name]+location_employees[name]+location_infra[name]))
 location_rows=[]
 for name in locations[:5]:location_rows.append((name,f"A {location_assets[name]} / E {location_employees[name]} / I {location_infra[name]}"))
 if len(locations)>5:
  rest=locations[5:];location_rows.append(("Other",f"A {sum(location_assets[x] for x in rest)} / E {sum(location_employees[x] for x in rest)} / I {sum(location_infra[x] for x in rest)}"))
 location_data=[[Paragraph("<b>Location-wise Summary</b>",styles["ArchiveCell"]),Paragraph("Assets / Employees / Infrastructure",styles["ArchiveCell"])]]+[[Paragraph(escape(name),styles["ArchiveCell"]),Paragraph(escape(value),styles["ArchiveCell"])] for name,value in location_rows]
 location_table=Table(location_data,colWidths=[180,280],repeatRows=1)
 location_table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1479b8")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#cbd5e1")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f5f8fb")])]))

 category_table=breakdown_table("Assets by Category",compact_breakdown(summary_frequency(assets,"Category"),4))
 department_table=breakdown_table("Assets by Department",compact_breakdown(summary_frequency(assets,"Department"),4))
 second_pair=Table([[category_table,department_table]],colWidths=[230,230]);second_pair.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),8)]))
 return [Paragraph("Executive Summary",styles["ArchiveSection"]),cards,Spacer(1,8),first_pair,Spacer(1,8),location_table,Spacer(1,8),second_pair]

def build_archive_pdf(path,start_date,end_date,sections):
 styles=getSampleStyleSheet()
 styles.add(ParagraphStyle(name="ArchiveTitle",parent=styles["Title"],textColor=colors.HexColor("#203f78"),alignment=TA_CENTER,spaceAfter=8))
 styles.add(ParagraphStyle(name="ArchiveSection",parent=styles["Heading1"],textColor=colors.HexColor("#1479b8"),spaceBefore=8,spaceAfter=8))
 styles.add(ParagraphStyle(name="ArchiveRecord",parent=styles["Heading3"],textColor=colors.HexColor("#203f78"),spaceBefore=7,spaceAfter=4))
 styles.add(ParagraphStyle(name="ArchiveCell",parent=styles["BodyText"],fontSize=7.5,leading=9.5))
 styles.add(ParagraphStyle(name="SummaryCard",parent=styles["BodyText"],alignment=TA_CENTER,textColor=colors.HexColor("#203f78"),leading=11))
 doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=36,leftMargin=36,topMargin=48,bottomMargin=34,title="ITAMS Data Archive")
 generated=datetime.now(LOCAL_TZ)
 story=[Paragraph("IT Asset Management",styles["ArchiveTitle"]),Paragraph("ITAMS Date-Range Data Archive",styles["Title"]),Paragraph(f"From {start_date:%d %b %Y} through {end_date:%d %b %Y}",styles["Heading2"]),Paragraph(f"Generated {generated:%d %b %Y %I:%M:%S %p} by {escape(session.get('user_name') or 'System')}",styles["Normal"]),Spacer(1,5),Paragraph("Dashboard summary and detailed records for the selected period. Password and session secrets are excluded. This PDF cannot restore SQLite.",styles["ArchiveCell"]),Spacer(1,6)]
 story.extend(archive_summary_flowables(sections,styles));story.append(PageBreak())
 for section_index,(title,rows) in enumerate(sections):
  story.append(Paragraph(f"{escape(title)} ({len(rows)})",styles["ArchiveSection"]))
  if not rows:story.append(Paragraph("No records changed in the selected date range.",styles["ArchiveCell"]))
  for index,row in enumerate(rows,1):
   populated=[(key,value) for key,value in row.items() if value not in (None,"")]
   identity=next((str(value) for key,value in populated if key in ("Location Asset Reference","Permanent System Reference","Asset Tag","Employee Code","Infrastructure Reference","Infrastructure Id","Username","Movement Id","Event Key")),f"Record {index}")
   story.append(Paragraph(f"{index}. {escape(identity)}",styles["ArchiveRecord"]))
   data=[[Paragraph(f"<b>{escape(str(key))}</b>",styles["ArchiveCell"]),Paragraph(escape(str(display_value(value))).replace("\n","<br/>"),styles["ArchiveCell"])] for key,value in populated]
   if data:
    table=Table(data,colWidths=[145,315],splitByRow=1)
    table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#eef4f8")),("LINEBELOW",(0,0),(-1,-1),0.2,colors.HexColor("#d8e0e8")),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3)]))
    story.append(table)
  if section_index<len(sections)-1:story.append(PageBreak())
 doc.build(story,onFirstPage=archive_page,onLaterPages=archive_page)

@bp.route("/archive/preview")
@login_required
def archive_preview():
 start_date,end_date,sections=archive_sections()
 return render_template("archive_preview.html",date_from=start_date,date_to=end_date,sections=[(title,len(rows)) for title,rows in sections],generated_at=datetime.now(LOCAL_TZ))

@bp.route("/archive/pdf")
@login_required
def archive_pdf():
 start_date,end_date,sections=archive_sections();stamp=datetime.now(LOCAL_TZ).strftime("%Y%m%d_%H%M%S")
 folder=Path(current_app.config["BACKUP_DIR"])/"pdf_archives";folder.mkdir(parents=True,exist_ok=True)
 path=folder/f"ITAMS_Data_Archive_{start_date:%Y%m%d}_{end_date:%Y%m%d}_{stamp}.pdf"
 build_archive_pdf(path,start_date,end_date,sections)
 audit("Report",path.name,"PDF Data Archive Export",new={"date_from":str(start_date),"date_to":str(end_date),"sections":{title:len(rows) for title,rows in sections}},reason="Authorized date-range archival export");db.session.commit()
 return send_file(path,as_attachment=True,download_name=path.name,mimetype="application/pdf")

def send_rows(rows,fmt,title,filename):
 if not rows:rows=[{"Result":"No records"}]
 def safe_cell(value):
  return "'"+value if isinstance(value,str) and value.startswith(("=","+","-","@")) else value
 columns,rows=rectangular_rows(rows)
 df=pd.DataFrame([{key:safe_cell(value) for key,value in row.items()} for row in rows],columns=columns);out=io.BytesIO()
 if fmt=="csv":out.write(df.to_csv(index=False).encode());mime="text/csv";suffix="csv"
 elif fmt=="excel":
  with pd.ExcelWriter(out,engine="openpyxl") as writer:
   df.to_excel(writer,index=False,sheet_name=title[:31],startrow=4)
   ws=writer.book[title[:31]];ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(6,max(1,len(df.columns))));ws.cell(1,1,"IT Asset Management")
   ws.cell(2,1,title);ws.cell(3,1,"Generated: "+datetime.now(LOCAL_TZ).strftime("%d %b %Y %I:%M:%S %p"));ws.cell(4,1,f"Records: {len(df)} | Populated fields exported: {len(df.columns)}")
   from openpyxl.styles import Alignment,Font,PatternFill
   from openpyxl.utils import get_column_letter
   ws.cell(1,1).font=Font(size=18,bold=True,color="203F78");ws.cell(2,1).font=Font(size=14,bold=True,color="76B82A")
   for cell in ws[5]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="203F78")
   ws.freeze_panes="A6";ws.auto_filter.ref=f"A5:{get_column_letter(len(df.columns))}{len(df)+5}";ws.sheet_view.showGridLines=False
   for index,column in enumerate(ws.iter_cols(),1):
    letter=get_column_letter(index);ws.column_dimensions[letter].width=min(45,max(12,max(len(str(c.value or "")) for c in column)+2))
    for cell in column[5:]:cell.alignment=Alignment(vertical="top",wrap_text=True)
  mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet";suffix="xlsx"
 elif fmt=="pdf":
  columns=list(df.columns);data=[columns]+[[str(x or "") for x in row] for row in df.itertuples(index=False,name=None)];styles=getSampleStyleSheet();table=Table(data,repeatRows=1);table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#203f78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("LINEBELOW",(0,0),(-1,0),2,colors.HexColor("#76b82a")),("FONTSIZE",(0,0),(-1,-1),7)]));SimpleDocTemplate(out,pagesize=A4).build(branded_pdf_heading(title,styles)+[table]);mime="application/pdf";suffix="pdf"
 else:return "Unsupported",400
 out.seek(0);return send_file(out,as_attachment=True,download_name=f"{filename}.{suffix}",mimetype=mime)

@bp.route("/dataset/<kind>/<fmt>")
@login_required
def dataset(kind,fmt):
 if fmt!="download":return redirect(url_for("reports.preview",kind=kind))
 return send_rows(dataset_rows(kind),"excel",kind.title(),kind)
