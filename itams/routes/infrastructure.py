from __future__ import annotations

from datetime import datetime, date, timedelta, timezone
from pathlib import Path
import io
import ipaddress
import json
import re
import uuid
from zoneinfo import ZoneInfo

import pandas as pd
from flask import (
    Blueprint, abort, current_app, flash, redirect, render_template, request,
    send_file, send_from_directory, session, url_for,
)
from sqlalchemy import func, or_, select
from werkzeug.utils import secure_filename

from ..extensions import db
from ..models import (
    AuditLog, Department, InfrastructureCategory, InfrastructureDocument,
    InfrastructureCable, InfrastructureHistory, InfrastructureIncident,
    InfrastructureItem, InfrastructureLocation, InfrastructurePort,
    InfrastructureQRIdentity, InfrastructureRelationship, InfrastructureVerification,
    Manufacturer, Vendor, Location,
)
from ..services.audit import audit
from ..services.lbac import (
    account_location_ids, can_access_location, is_super_admin,
    validate_location, workspace_locations,
)
from ..services.master_values import resolve_master_id
from ..services.uploads import validate_upload
from ..services.infrastructure import (
    EVENT_TYPES, active_identity, all_specification_fields, ensure_infrastructure_qr,
    infrastructure_profile_url, infrastructure_qr_path, qr_history, qr_metadata,
    regenerate_infrastructure_qr, resolve_infrastructure_token, serialize_details,
    specification_fields, specifications_from_form,
)
from .helpers import can, login_required, normalized_role


bp = Blueprint("infrastructure", __name__, url_prefix="/infrastructure")
ALLOWED_DOCUMENT_EXTENSIONS = {
    "pdf", "png", "jpg", "jpeg", "webp", "doc", "docx", "xls", "xlsx", "csv", "txt",
}
STATUS_OPTIONS = ["Planned", "Installed", "Active", "Under Maintenance", "Under Repair", "Inactive", "Decommissioned"]
LOCATION_TYPES = ["Site", "Building", "Floor", "Department", "Room", "Rack", "Rack Unit"]
RELATIONSHIP_TYPES = ["Connected To", "Uplink To", "Downlink To", "Managed By", "Protected By", "Feeds", "Member Of", "Other"]
CONNECTION_STATUSES = ["Connected", "Disconnected", "Planned", "Faulty", "Unknown"]
PORT_STATUSES = ["Available", "Connected", "Disabled", "Faulty", "Reserved"]
PORT_CATEGORIES = {"Routers", "Switches", "Patch Panels", "EPABX", "NVR/DVR"}
NETWORK_CATEGORIES = {
    "Routers", "Switches", "Firewalls", "Wireless Access Points", "Cisco Access Points", "Modems",
    "EPABX", "VoIP Phones", "Biometric Devices", "NVR/DVR",
    "Printers", "Scanners", "Copiers",
}
CABLE_TYPES = ["Cat5e", "Cat6", "Cat6A", "Fiber", "Telephone", "HDMI", "USB", "Display Cable"]
CABLE_STATUSES = ["Installed", "Available", "Disconnected", "Faulty", "Retired"]
INCIDENT_STATUSES = ["Open", "In Progress", "On Hold", "Resolved", "Closed"]
INCIDENT_PRIORITIES = ["Low", "Medium", "High", "Critical"]
PHOTO_TYPES = ["Front", "Back", "Installation", "Damage", "Serial Label", "Other"]


@bp.before_request
def restrict_infrastructure():
    if not normalized_role():
        return redirect(url_for("auth.login", next=request.full_path if request.query_string else request.path))
    if request.endpoint in {
        "infrastructure.locations", "infrastructure.location_toggle",
        "infrastructure.cables", "infrastructure.update_cable", "infrastructure.delete_cable",
    }:
        abort(404)
    permissions = {
        "infrastructure.index": "infrastructure.view",
        "infrastructure.profile": "infrastructure.view",
        "infrastructure.qr_resolve": "infrastructure.view",
        "infrastructure.qr_image": "infrastructure.view",
        "infrastructure.qr_download": "infrastructure.view",
        "infrastructure.qr_label": "infrastructure.view",
        "infrastructure.document": "infrastructure.view",
        "infrastructure.new": "infrastructure.add",
        "infrastructure.edit": "infrastructure.edit",
        "infrastructure.add_history": "infrastructure.history.manage",
        "infrastructure.upload_document": "infrastructure.document.upload",
        "infrastructure.delete": "infrastructure.delete",
        "infrastructure.export_excel": "infrastructure.export",
        "infrastructure.qr_regenerate": "infrastructure.qr.manage",
        "infrastructure.qr_manage": "infrastructure.qr.manage",
        "infrastructure.locations": "infrastructure.view",
        "infrastructure.location_toggle": "infrastructure.location.manage",
        "infrastructure.add_relationship": "infrastructure.relationship.manage",
        "infrastructure.update_relationship": "infrastructure.relationship.manage",
        "infrastructure.delete_relationship": "infrastructure.relationship.manage",
        "infrastructure.add_port": "infrastructure.port.manage",
        "infrastructure.update_port": "infrastructure.port.manage",
        "infrastructure.delete_port": "infrastructure.port.manage",
        "infrastructure.cables": "infrastructure.view",
        "infrastructure.update_cable": "infrastructure.cable.manage",
        "infrastructure.delete_cable": "infrastructure.cable.manage",
        "infrastructure.add_incident": "infrastructure.incident.manage",
        "infrastructure.update_incident": "infrastructure.incident.manage",
        "infrastructure.add_verification": "infrastructure.verification.manage",
        "infrastructure.reports": "infrastructure.report.view",
    }
    if request.endpoint == "infrastructure.locations" and request.method == "POST":
        permissions["infrastructure.locations"] = "infrastructure.location.manage"
    if request.endpoint == "infrastructure.cables" and request.method == "POST":
        permissions["infrastructure.cables"] = "infrastructure.cable.manage"
    if request.endpoint == "infrastructure.reports" and request.args.get("format") == "excel":
        permissions["infrastructure.reports"] = "infrastructure.export"
    required = permissions.get(request.endpoint, "infrastructure.view")
    if not can(required):
        abort(403)


def references():
    return {
        "categories": InfrastructureCategory.query.filter_by(active=True).order_by(InfrastructureCategory.name).all(),
        "manufacturers": Manufacturer.query.filter_by(active=True).order_by(Manufacturer.name).all(),
        "vendors": Vendor.query.filter_by(active=True).order_by(Vendor.name).all(),
        "departments": Department.query.filter_by(active=True).order_by(Department.name).all(),
        "status_options": STATUS_OPTIONS,
        "specification_fields": all_specification_fields(),
        "event_types": EVENT_TYPES,
        "locations": InfrastructureLocation.query.filter_by(active=True).order_by(InfrastructureLocation.location_type, InfrastructureLocation.name).all(),
        # Dashboard workspace selection must not hide other locations assigned
        # to this account from the infrastructure form.
        "company_locations": workspace_locations(),
        "infrastructure_items": InfrastructureItem.query.order_by(InfrastructureItem.infrastructure_reference).all(),
        "location_types": LOCATION_TYPES, "relationship_types": RELATIONSHIP_TYPES,
        "connection_statuses": CONNECTION_STATUSES, "port_statuses": PORT_STATUSES, "port_categories": PORT_CATEGORIES,
        "network_categories": NETWORK_CATEGORIES, "cable_types": CABLE_TYPES, "cable_statuses": CABLE_STATUSES,
        "incident_statuses": INCIDENT_STATUSES, "incident_priorities": INCIDENT_PRIORITIES,
        "photo_types": PHOTO_TYPES,
    }


def parse_date(value):
    return datetime.strptime(value, "%Y-%m-%d").date() if value else None


def optional_ip(value, label):
    value = (value or "").strip()
    if not value:
        return None
    try:
        ipaddress.ip_address(value)
        return value
    except ValueError as exc:
        if label == "Management IP":
            host_port = re.fullmatch(r"(\d{1,3}(?:\.\d{1,3}){3}):(\d{1,5})", value)
            if host_port:
                host, port = host_port.groups()
                try:
                    ipaddress.ip_address(host)
                except ValueError:
                    pass
                else:
                    if 1 <= int(port) <= 65535:
                        return host
        raise ValueError(f"{label} is not a valid IP address. Use a documentation-range value such as 192.0.2.10; IP:port is also accepted.") from exc


def resolve_company_location(form):
    """Resolve an optional location using account scope, not dashboard workspace scope."""
    raw = (form.get("location_id") or "").strip()
    if not raw:
        ids = account_location_ids()
        if ids is not None:
            if len(ids) != 1:
                raise ValueError("Select one of the Locations assigned to your account.")
            return db.session.execute(
                select(Location).where(Location.id == ids[0], Location.active == True).execution_options(lbac_bypass=True)
            ).scalar_one_or_none()
        return None
    if raw == "__other__":
        name = (form.get("location_id_other") or "").strip()
        if not name:
            raise ValueError("Enter the new Location.")
        location = db.session.execute(
            select(Location)
            .where(func.lower(func.trim(Location.name)) == name.lower())
            .execution_options(lbac_bypass=True)
        ).scalar_one_or_none()
        if location and not location.active:
            raise ValueError(f"Location '{name}' exists but is inactive. Activate it from Administration.")
        if not location:
            if not is_super_admin():
                raise ValueError("Only Super Admin can create a new company location. Select an assigned location or Not specified.")
            base = re.sub(r"[^A-Z0-9]+", "", name.upper())[:12] or "LOC"
            code = base
            suffix = 2
            while db.session.execute(
                select(Location.id)
                .where(func.lower(Location.code) == code.lower())
                .execution_options(lbac_bypass=True)
            ).first():
                code = f"{base[:max(1, 12-len(str(suffix)))]}{suffix}"
                suffix += 1
            location = Location(name=name, code=code, active=True)
            db.session.add(location)
            db.session.flush()
            audit("Location", location.id, "Created from Infrastructure Other option", new={"name": name, "code": code})
    else:
        try:
            location_id = int(raw)
        except (TypeError, ValueError) as exc:
            raise ValueError("Select a valid Location.") from exc
        location = db.session.execute(
            select(Location)
            .where(Location.id == location_id, Location.active == True)
            .execution_options(lbac_bypass=True)
        ).scalar_one_or_none()
        if not location:
            raise ValueError("Select an active Location.")
    ids = account_location_ids()
    if ids is not None and location.id not in ids:
        raise ValueError("Select a Location assigned to your account.")
    return location


def validate_network(item):
    effective = item.management_ip or item.ip_address
    if effective:
        duplicate = InfrastructureItem.query.filter(
            InfrastructureItem.id != (item.id or 0),
            func.lower(func.trim(func.coalesce(InfrastructureItem.management_ip, InfrastructureItem.ip_address))) == effective.lower(),
        ).first()
        if duplicate:
            raise ValueError(f"Management IP {effective} is already registered to {duplicate.infrastructure_reference}.")
    if item.subnet:
        try:
            ipaddress.ip_network(item.subnet, strict=False)
        except ValueError as exc:
            raise ValueError("Subnet must be a valid network or CIDR value.") from exc
    if item.mac_address and not re.fullmatch(r"[0-9A-Fa-f]{2}([:\-])[0-9A-Fa-f]{2}(?:\1[0-9A-Fa-f]{2}){4}", item.mac_address):
        raise ValueError("MAC Address must use six hexadecimal pairs separated by colons or hyphens.")


def populate(item, form):
    category_id = resolve_master_id(
        InfrastructureCategory, form, "category_id",
        required=True, label="Infrastructure Category",
    )
    category = db.session.get(InfrastructureCategory, category_id)
    item.category = category
    location = resolve_company_location(form)
    item.location_id = location.id if location else None
    for field in (
        "name", "model", "serial_number", "asset_tag", "site", "floor", "room",
        "condition", "invoice_number", "mac_address", "remarks", "dns", "network_vlan",
    ):
        setattr(item, field, (form.get(field) or "").strip() or None)
    item.manufacturer_id = resolve_master_id(Manufacturer, form, "manufacturer_id", label="Manufacturer")
    item.vendor_id = resolve_master_id(Vendor, form, "vendor_id", label="Vendor")
    item.department_id = resolve_master_id(Department, form, "department_id", label="Department")
    # Update only the currently active category fields. Retired fields remain
    # stored in JSON so hiding a field never destroys historical information.
    specifications = dict(item.specifications)
    submitted_specifications = specifications_from_form(category.name, form)
    for code, _label in specification_fields(category.name):
        if code in submitted_specifications:
            specifications[code] = submitted_specifications[code]
        else:
            specifications.pop(code, None)
    raw_management_ip = (form.get("management_ip") or form.get("ip_address") or "").strip()
    try:
        item.management_ip = optional_ip(raw_management_ip, "Management IP")
    except ValueError:
        if category.name != "CCTV Cameras":
            raise
        # CCTV Network Parameter deliberately accepts endpoints, ports, VLANs
        # and other free-form connection details. It must not block an edit.
        if raw_management_ip and not specifications.get("network_parameter"):
            specifications["network_parameter"] = raw_management_ip
        item.management_ip = None
    item.ip_address = item.management_ip
    item.gateway = optional_ip(form.get("gateway"), "Gateway")
    item.subnet = (form.get("subnet") or "").strip() or None
    item.ip_assignment = form.get("ip_assignment") if form.get("ip_assignment") in ("DHCP", "Static") else None
    item.status = form.get("status") if form.get("status") in STATUS_OPTIONS else "Planned"
    for field in ("purchase_date", "installation_date", "warranty_expiry", "amc_start_date", "amc_expiry"):
        setattr(item, field, parse_date(form.get(field)))
    item.purchase_cost = form.get("purchase_cost", type=float)
    item.specifications_json = json.dumps(specifications)
    validate_network(item)
    item.updated_by = session.get("user_name")
    item.state_version = (item.state_version or 0) + 1


def filtered_query():
    query = InfrastructureItem.query
    q = request.args.get("q", "").strip()
    if q:
        query = query.outerjoin(InfrastructureItem.category).outerjoin(InfrastructureItem.manufacturer).outerjoin(InfrastructureItem.vendor).outerjoin(InfrastructureItem.department).outerjoin(InfrastructureItem.location_node).filter(or_(
            InfrastructureItem.infrastructure_reference.contains(q),
            InfrastructureItem.name.contains(q), InfrastructureItem.model.contains(q),
            InfrastructureItem.serial_number.contains(q), InfrastructureItem.asset_tag.contains(q),
            InfrastructureItem.ip_address.contains(q), InfrastructureItem.management_ip.contains(q),
            InfrastructureItem.mac_address.contains(q), InfrastructureItem.subnet.contains(q),
            InfrastructureItem.gateway.contains(q), InfrastructureItem.dns.contains(q),
            InfrastructureItem.network_vlan.contains(q), InfrastructureItem.site.contains(q),
            InfrastructureItem.floor.contains(q), InfrastructureItem.room.contains(q),
            InfrastructureCategory.name.contains(q),
            Manufacturer.name.contains(q), Vendor.name.contains(q), Department.name.contains(q),
            InfrastructureItem.specifications_json.contains(q),
        ))
    for key, column in (
        ("category", InfrastructureItem.category_id), ("department", InfrastructureItem.department_id),
        ("manufacturer", InfrastructureItem.manufacturer_id), ("vendor", InfrastructureItem.vendor_id),
        ("status", InfrastructureItem.status), ("location", InfrastructureItem.location_id),
    ):
        if request.args.get(key):
            query = query.filter(column == request.args[key])
    for key, column in (
        ("site", InfrastructureItem.site), ("floor", InfrastructureItem.floor),
        ("room", InfrastructureItem.room), ("management_ip", InfrastructureItem.management_ip),
        ("firmware", InfrastructureItem.specifications_json),
    ):
        if request.args.get(key):
            query = query.filter(column.contains(request.args[key].strip()))
    warranty = request.args.get("warranty")
    if warranty == "expired":
        query = query.filter(InfrastructureItem.warranty_expiry < date.today())
    elif warranty == "expiring":
        query = query.filter(InfrastructureItem.warranty_expiry.between(date.today(), date.today() + timedelta(days=30)))
    elif warranty == "valid":
        query = query.filter(InfrastructureItem.warranty_expiry >= date.today())
    amc = request.args.get("amc")
    if amc == "expired":
        query = query.filter(InfrastructureItem.amc_expiry < date.today())
    elif amc == "expiring":
        query = query.filter(InfrastructureItem.amc_expiry.between(date.today(), date.today() + timedelta(days=30)))
    elif amc == "valid":
        query = query.filter(InfrastructureItem.amc_expiry >= date.today())
    return query


@bp.route("/")
@login_required
def index():
    page = max(1, request.args.get("page", 1, type=int))
    pagination = filtered_query().order_by(InfrastructureItem.infrastructure_reference).paginate(page=page, per_page=50, error_out=False)
    return render_template("infrastructure_list.html", items=pagination.items, pagination=pagination, **references())


@bp.route("/new", methods=["GET", "POST"])
@login_required
def new():
    if request.method == "POST":
        try:
            item = InfrastructureItem(status="Planned")
            populate(item, request.form)
            db.session.add(item)
            db.session.flush()
            db.session.add(InfrastructureHistory(
                infrastructure=item, event_type="Record Created", performed_by=session.get("user_name", "System"),
                reason="Infrastructure master created", remarks=item.remarks,
            ))
            audit("Infrastructure", item.id, "Create", new={
                "reference": item.infrastructure_reference, "category": item.category.name,
            })
            ensure_infrastructure_qr(item, actor=session.get("user_name"))
            db.session.commit()
            if is_super_admin():
                if item.location_id:
                    session["workspace_location_id"] = item.location_id
                else:
                    session.pop("workspace_location_id", None)
            flash("Infrastructure record created with permanent ID and QR code.", "success")
            return redirect(url_for("infrastructure.profile", reference=item.infrastructure_reference))
        except (ValueError, RuntimeError) as exc:
            db.session.rollback()
            flash(str(exc), "danger")
        except Exception:
            db.session.rollback()
            current_app.logger.exception("Could not create infrastructure record")
            flash("Could not create the infrastructure record.", "danger")
    return render_template("infrastructure_form.html", item=None, **references())


@bp.route("/profile/<string:reference>")
@login_required
def profile(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    qr_error = None
    try:
        _path, _url, created = ensure_infrastructure_qr(item)
        if created:
            db.session.commit()
    except RuntimeError as exc:
        db.session.rollback()
        qr_error = str(exc)
    return render_template(
        "infrastructure_profile.html", item=item, qr_error=qr_error,
        profile_url=infrastructure_profile_url(item) if not qr_error else None,
        qr_meta=qr_metadata(item), qr_history_rows=qr_history(item.id),
        history=InfrastructureHistory.query.filter_by(infrastructure_id=item.id).order_by(InfrastructureHistory.occurred_at.desc()).all(),
        documents=InfrastructureDocument.query.filter_by(infrastructure_id=item.id).order_by(InfrastructureDocument.created_at.desc()).all(),
        relationships_out=InfrastructureRelationship.query.filter_by(source_id=item.id).order_by(InfrastructureRelationship.relationship_type).all(),
        relationships_in=InfrastructureRelationship.query.filter_by(target_id=item.id).order_by(InfrastructureRelationship.relationship_type).all(),
        ports=InfrastructurePort.query.filter_by(infrastructure_id=item.id).order_by(InfrastructurePort.port_number).all(),
        verifications=InfrastructureVerification.query.filter_by(infrastructure_id=item.id).order_by(InfrastructureVerification.verified_date.desc()).all(),
        audit_rows=AuditLog.query.filter_by(entity="Infrastructure", entity_id=str(item.id)).order_by(AuditLog.timestamp.desc()).all(),
        category_specification_fields=specification_fields(item.category.name),
        **references(),
    )


@bp.route("/profile/<string:reference>/edit", methods=["GET", "POST"])
@login_required
def edit(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    if request.method == "POST":
        reason = request.form.get("reason", "").strip()
        if not reason:
            flash("Reason for edit is required.", "danger")
        else:
            try:
                old = {
                    "category": item.category.name, "status": item.status,
                    "site": item.site, "building": item.building, "floor": item.floor, "room": item.room,
                    "rack": item.rack, "rack_unit": item.rack_unit,
                    "management_ip": item.effective_management_ip, "mac_address": item.mac_address,
                    "specifications": item.specifications,
                }
                populate(item, request.form)
                new_values = {
                    "category": item.category.name, "status": item.status,
                    "site": item.site, "building": item.building, "floor": item.floor, "room": item.room,
                    "rack": item.rack, "rack_unit": item.rack_unit,
                    "management_ip": item.effective_management_ip, "mac_address": item.mac_address,
                    "specifications": item.specifications,
                }
                audit("Infrastructure", item.id, "Update", old=old, new=new_values, reason=reason)
                if old["specifications"] != new_values["specifications"]:
                    audit("Infrastructure", item.id, "Specification Update", old=old["specifications"], new=new_values["specifications"], reason=reason)
                db.session.commit()
                if is_super_admin():
                    if item.location_id:
                        session["workspace_location_id"] = item.location_id
                    else:
                        session.pop("workspace_location_id", None)
                flash("Infrastructure record updated.", "success")
                return redirect(url_for("infrastructure.profile", reference=reference))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "danger")
    return render_template("infrastructure_form.html", item=item, **references())


@bp.route("/profile/<string:reference>/history", methods=["POST"])
@login_required
def add_history(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    event_type = request.form.get("event_type", "")
    if event_type not in EVENT_TYPES:
        abort(400)
    reason = request.form.get("reason", "").strip()
    occurred = parse_date(request.form.get("occurred_date"))
    now_utc = datetime.now(timezone.utc).replace(tzinfo=None)
    occurred_at = datetime.combine(occurred, now_utc.time()) if occurred else now_utc
    details = {
        "vendor": request.form.get("event_vendor"), "cost": request.form.get("event_cost"),
        "ticket_number": request.form.get("ticket_number"),
        "firmware_before": request.form.get("firmware_before"),
        "firmware_after": request.form.get("firmware_after"),
        "from_location": request.form.get("from_location"),
        "to_location": request.form.get("to_location"),
        "location_hierarchy": request.form.get("new_location_node_id"),
        "next_due_date": request.form.get("next_due_date"),
        "outcome": request.form.get("outcome"),
    }
    history = InfrastructureHistory(
        infrastructure=item, event_type=event_type, occurred_at=occurred_at,
        reason=reason or None, remarks=request.form.get("remarks", "").strip() or None,
        details_json=serialize_details(details), performed_by=session.get("user_name", "System"),
    )
    db.session.add(history)
    if event_type == "Installation":
        item.status = "Installed"
        item.installation_date = occurred or item.installation_date
    elif event_type == "Maintenance":
        item.status = "Under Maintenance" if request.form.get("outcome") != "Completed" else "Active"
    elif event_type == "Repair":
        item.status = "Under Repair" if request.form.get("outcome") != "Completed" else "Active"
    elif event_type == "Firmware Update" and request.form.get("firmware_after"):
        specs = item.specifications
        specs["firmware_version"] = request.form["firmware_after"].strip()
        item.specifications_json = json.dumps(specs)
    elif event_type == "Relocation":
        item.site = request.form.get("new_site", "").strip() or item.site
        item.floor = request.form.get("new_floor", "").strip() or item.floor
        item.room = request.form.get("new_room", "").strip() or item.room
        item.department_id = request.form.get("new_department_id", type=int) or item.department_id
    elif event_type == "Warranty" and request.form.get("warranty_expiry"):
        item.warranty_expiry = parse_date(request.form["warranty_expiry"])
    elif event_type == "AMC":
        item.amc_start_date = parse_date(request.form.get("amc_start_date")) or item.amc_start_date
        item.amc_expiry = parse_date(request.form.get("amc_expiry")) or item.amc_expiry
    elif event_type == "Decommissioning":
        item.status = "Decommissioned"
    item.updated_by = session.get("user_name")
    item.state_version = (item.state_version or 0) + 1
    audit("Infrastructure", item.id, event_type, new=details, reason=reason)
    db.session.commit()
    flash(f"{event_type} history recorded.", "success")
    return redirect(url_for("infrastructure.profile", reference=reference))


@bp.route("/profile/<string:reference>/documents", methods=["POST"])
@login_required
def upload_document(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    upload = request.files.get("document")
    try:
        original = validate_upload(upload)
    except ValueError as exc:
        flash(str(exc), "danger")
        return redirect(url_for("infrastructure.profile", reference=reference))
    extension = original.rsplit(".", 1)[-1].lower()
    folder = Path(current_app.config["DOCUMENT_DIR"]) / "infrastructure_documents"
    folder.mkdir(parents=True, exist_ok=True)
    stored = f"{uuid.uuid4().hex}_{original}"
    target = folder / stored
    upload.save(target)
    document = InfrastructureDocument(
        infrastructure=item, original_name=original, stored_name=stored,
        content_type=upload.mimetype, size_bytes=target.stat().st_size,
        is_photo=extension in {"png", "jpg", "jpeg", "webp"},
        photo_type=request.form.get("photo_type") if request.form.get("photo_type") in PHOTO_TYPES else None,
        uploaded_by=session.get("user_name"),
    )
    db.session.add(document)
    db.session.flush()
    audit("Infrastructure", item.id, "Document Uploaded", new={"name": original, "photo": document.is_photo, "photo_type": document.photo_type})
    db.session.commit()
    flash("Document uploaded.", "success")
    return redirect(url_for("infrastructure.profile", reference=reference))


@bp.route("/documents/<int:document_id>")
@login_required
def document(document_id):
    document = db.session.get(InfrastructureDocument, document_id)
    if not document:
        abort(404)
    folder = Path(current_app.config["DOCUMENT_DIR"]) / "infrastructure_documents"
    if not (folder / document.stored_name).exists():
        folder = Path(current_app.config["UPLOAD_DIR"]) / "infrastructure_documents"
    return send_from_directory(
        folder,
        document.stored_name, as_attachment=not document.is_photo,
        download_name=document.original_name,
    )


@bp.route("/locations", methods=["GET", "POST"])
@login_required
def locations():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        location_type = request.form.get("location_type", "")
        parent_id = request.form.get("parent_id", type=int)
        parent = db.session.get(InfrastructureLocation, parent_id) if parent_id else None
        company_location = validate_location(request.form.get("company_location_id", type=int) or (parent.company_location_id if parent else None))
        if not name or location_type not in LOCATION_TYPES:
            flash("Location name and a valid hierarchy level are required.", "danger")
        elif parent and LOCATION_TYPES.index(parent.location_type) >= LOCATION_TYPES.index(location_type):
            flash("The parent location must be above the selected hierarchy level.", "danger")
        elif parent and parent.company_location_id != company_location.id:
            flash("Parent and child hierarchy nodes must use the same company location.", "danger")
        else:
            duplicate = InfrastructureLocation.query.filter(
                InfrastructureLocation.parent_id == (parent.id if parent else None),
                func.lower(InfrastructureLocation.name) == name.lower(),
                InfrastructureLocation.location_type == location_type,
            ).first()
            if duplicate:
                flash("That location already exists under the selected parent.", "danger")
            else:
                row = InfrastructureLocation(name=name, location_type=location_type, parent=parent, company_location=company_location, created_by=session.get("user_name"))
                db.session.add(row);db.session.flush()
                audit("Infrastructure Location", row.id, "Create", new={"name": name, "type": location_type, "parent": parent.full_path if parent else None},location_id=company_location.id)
                db.session.commit();flash("Infrastructure location added.", "success")
        return redirect(url_for("infrastructure.locations"))
    rows = InfrastructureLocation.query.order_by(InfrastructureLocation.location_type, InfrastructureLocation.name).all()
    return render_template("infrastructure_locations.html", rows=rows, **references())


@bp.route("/locations/<int:location_id>/toggle", methods=["POST"])
@login_required
def location_toggle(location_id):
    row = db.session.get(InfrastructureLocation, location_id)
    if not row:abort(404)
    if row.active and (row.children.filter_by(active=True).count() or row.items):
        flash("Cannot deactivate a location that has active child locations or assigned infrastructure items.","danger")
        return redirect(url_for("infrastructure.locations"))
    old=row.active;row.active=not row.active
    audit("Infrastructure Location",row.id,"Activate" if row.active else "Deactivate",old={"active":old},new={"active":row.active},reason=request.form.get("reason"))
    db.session.commit();flash("Location status updated.","success")
    return redirect(url_for("infrastructure.locations"))


@bp.route("/profile/<string:reference>/relationships", methods=["POST"])
@login_required
def add_relationship(reference):
    item=InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    target=db.session.get(InfrastructureItem,request.form.get("target_id",type=int))
    relationship_type=request.form.get("relationship_type","")
    if not item or not target:abort(404)
    if item.id==target.id:
        flash("A device cannot be connected to itself.","danger")
    elif relationship_type not in RELATIONSHIP_TYPES:
        flash("Select a valid relationship type.","danger")
    elif InfrastructureRelationship.query.filter_by(source_id=item.id,target_id=target.id,relationship_type=relationship_type).first():
        flash("This device relationship already exists.","danger")
    else:
        row=InfrastructureRelationship(source=item,target=target,relationship_type=relationship_type,connection_status=request.form.get("connection_status") if request.form.get("connection_status") in CONNECTION_STATUSES else "Connected",remarks=request.form.get("remarks","").strip() or None,created_by=session.get("user_name"))
        db.session.add(row);db.session.flush()
        audit("Infrastructure",item.id,"Relationship Added",new={"target":target.infrastructure_reference,"type":relationship_type,"status":row.connection_status},reason=row.remarks)
        audit("Infrastructure",target.id,"Relationship Added",new={"source":item.infrastructure_reference,"type":relationship_type,"status":row.connection_status},reason=row.remarks)
        db.session.commit();flash("Device relationship added.","success")
    return redirect(url_for("infrastructure.profile",reference=reference))


@bp.route("/relationships/<int:relationship_id>/update", methods=["POST"])
@login_required
def update_relationship(relationship_id):
    row=db.session.get(InfrastructureRelationship,relationship_id)
    if not row:abort(404)
    status=request.form.get("connection_status")
    if status not in CONNECTION_STATUSES:abort(400)
    old={"status":row.connection_status,"remarks":row.remarks}
    row.connection_status=status;row.remarks=request.form.get("remarks","").strip() or None;row.updated_by=session.get("user_name")
    details={"source":row.source.infrastructure_reference,"target":row.target.infrastructure_reference,"type":row.relationship_type,"status":status,"remarks":row.remarks}
    audit("Infrastructure",row.source_id,"Relationship Updated",old=old,new=details,reason=request.form.get("reason"))
    audit("Infrastructure",row.target_id,"Relationship Updated",old=old,new=details,reason=request.form.get("reason"))
    db.session.commit();flash("Device relationship updated.","success")
    return redirect(url_for("infrastructure.profile",reference=row.source.infrastructure_reference))


@bp.route("/relationships/<int:relationship_id>/delete", methods=["POST"])
@login_required
def delete_relationship(relationship_id):
    row=db.session.get(InfrastructureRelationship,relationship_id)
    if not row:abort(404)
    reference=row.source.infrastructure_reference
    details={"source":row.source.infrastructure_reference,"target":row.target.infrastructure_reference,"type":row.relationship_type}
    audit("Infrastructure",row.source_id,"Relationship Removed",old=details,reason=request.form.get("reason"))
    audit("Infrastructure",row.target_id,"Relationship Removed",old=details,reason=request.form.get("reason"))
    db.session.delete(row);db.session.commit();flash("Device relationship removed.","success")
    return redirect(url_for("infrastructure.profile",reference=reference))


@bp.route("/profile/<string:reference>/ports", methods=["POST"])
@login_required
def add_port(reference):
    item=InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:abort(404)
    if item.category.name not in PORT_CATEGORIES:
        abort(400)
    port_number=request.form.get("port_number","").strip()
    connected=db.session.get(InfrastructureItem,request.form.get("connected_device_id",type=int))
    if not port_number:
        flash("Port Number is required.","danger")
    elif InfrastructurePort.query.filter_by(infrastructure_id=item.id,port_number=port_number).first():
        flash(f"Port {port_number} already exists on this device.","danger")
    elif connected and connected.id==item.id:
        flash("A port cannot connect to its own device.","danger")
    else:
        poe_value=request.form.get("poe")
        port_status=request.form.get("status","Available")
        if port_status not in PORT_STATUSES:abort(400)
        row=InfrastructurePort(infrastructure=item,port_number=port_number,port_name=request.form.get("port_name","").strip() or None,connected_device=connected,connected_port=request.form.get("connected_port","").strip() or None,status=port_status,speed=request.form.get("speed","").strip() or None,poe=True if poe_value=="Yes" else False if poe_value=="No" else None,vlan=request.form.get("vlan","").strip() or None,remarks=request.form.get("remarks","").strip() or None)
        db.session.add(row);db.session.flush()
        audit("Infrastructure",item.id,"Port Added",new={"port":port_number,"connected_device":connected.infrastructure_reference if connected else None,"connected_port":row.connected_port,"status":row.status})
        db.session.commit();flash("Network port added.","success")
    return redirect(url_for("infrastructure.profile",reference=reference))


@bp.route("/ports/<int:port_id>/update", methods=["POST"])
@login_required
def update_port(port_id):
    row=db.session.get(InfrastructurePort,port_id)
    if not row:abort(404)
    connected_id=request.form.get("connected_device_id",type=int)
    connected=db.session.get(InfrastructureItem,connected_id) if connected_id else None
    if connected and connected.id==row.infrastructure_id:
        flash("A port cannot connect to its own device.","danger")
        return redirect(url_for("infrastructure.profile",reference=row.infrastructure.infrastructure_reference))
    old={"status":row.status,"connected_device":row.connected_device.infrastructure_reference if row.connected_device else None,"connected_port":row.connected_port,"vlan":row.vlan}
    port_status=request.form.get("status","Available")
    if port_status not in PORT_STATUSES:abort(400)
    row.port_name=request.form.get("port_name","").strip() or None;row.connected_device=connected;row.connected_port=request.form.get("connected_port","").strip() or None;row.status=port_status;row.speed=request.form.get("speed","").strip() or None
    poe_value=request.form.get("poe");row.poe=True if poe_value=="Yes" else False if poe_value=="No" else None
    row.vlan=request.form.get("vlan","").strip() or None;row.remarks=request.form.get("remarks","").strip() or None
    audit("Infrastructure",row.infrastructure_id,"Port Updated",old=old,new={"port":row.port_number,"status":row.status,"connected_device":connected.infrastructure_reference if connected else None,"connected_port":row.connected_port,"vlan":row.vlan},reason=request.form.get("reason"))
    db.session.commit();flash("Network port updated.","success")
    return redirect(url_for("infrastructure.profile",reference=row.infrastructure.infrastructure_reference))


@bp.route("/ports/<int:port_id>/delete", methods=["POST"])
@login_required
def delete_port(port_id):
    row=db.session.get(InfrastructurePort,port_id)
    if not row:abort(404)
    reference=row.infrastructure.infrastructure_reference
    audit("Infrastructure",row.infrastructure_id,"Port Removed",old={"port":row.port_number,"connected_device":row.connected_device.infrastructure_reference if row.connected_device else None},reason=request.form.get("reason"))
    db.session.delete(row);db.session.commit();flash("Network port removed.","success")
    return redirect(url_for("infrastructure.profile",reference=reference))


@bp.route("/cables", methods=["GET", "POST"])
@login_required
def cables():
    if request.method=="POST":
        cable_id=request.form.get("cable_id","").strip()
        source=db.session.get(InfrastructureItem,request.form.get("source_device_id",type=int))
        destination=db.session.get(InfrastructureItem,request.form.get("destination_device_id",type=int))
        cable_type=request.form.get("cable_type","")
        if not cable_id or cable_type not in CABLE_TYPES:
            flash("Cable ID and a valid cable type are required.","danger")
        elif InfrastructureCable.query.filter(func.lower(InfrastructureCable.cable_id)==cable_id.lower()).first():
            flash("Cable ID already exists.","danger")
        elif source and destination and source.id==destination.id:
            flash("Cable source and destination devices must be different.","danger")
        else:
            cable_status=request.form.get("status","Installed")
            if cable_status not in CABLE_STATUSES:abort(400)
            row=InfrastructureCable(cable_id=cable_id,cable_type=cable_type,length=request.form.get("length","").strip() or None,source_device=source,destination_device=destination,source_port=request.form.get("source_port","").strip() or None,destination_port=request.form.get("destination_port","").strip() or None,installation_date=parse_date(request.form.get("installation_date")),status=cable_status,remarks=request.form.get("remarks","").strip() or None,created_by=session.get("user_name"))
            db.session.add(row);db.session.flush()
            details={"cable_id":cable_id,"type":cable_type,"source":source.infrastructure_reference if source else None,"destination":destination.infrastructure_reference if destination else None}
            audit("Infrastructure Cable",row.id,"Create",new=details)
            if source:audit("Infrastructure",source.id,"Cable Added",new=details)
            if destination:audit("Infrastructure",destination.id,"Cable Added",new=details)
            db.session.commit();flash("Cable record added.","success")
        return redirect(url_for("infrastructure.cables"))
    return render_template("infrastructure_cables.html",rows=InfrastructureCable.query.order_by(InfrastructureCable.cable_id).all(),**references())


@bp.route("/cables/<int:cable_id>/update", methods=["POST"])
@login_required
def update_cable(cable_id):
    row=db.session.get(InfrastructureCable,cable_id)
    if not row:abort(404)
    old={"status":row.status,"remarks":row.remarks}
    cable_status=request.form.get("status","Installed")
    if cable_status not in CABLE_STATUSES:abort(400)
    row.status=cable_status;row.remarks=request.form.get("remarks","").strip() or None
    details={"cable_id":row.cable_id,"status":row.status,"remarks":row.remarks}
    audit("Infrastructure Cable",row.id,"Update",old=old,new=details,reason=request.form.get("reason"))
    if row.source_device_id:audit("Infrastructure",row.source_device_id,"Cable Updated",old=old,new=details,reason=request.form.get("reason"))
    if row.destination_device_id:audit("Infrastructure",row.destination_device_id,"Cable Updated",old=old,new=details,reason=request.form.get("reason"))
    db.session.commit();flash("Cable status updated.","success")
    return redirect(url_for("infrastructure.cables"))


@bp.route("/cables/<int:cable_id>/delete", methods=["POST"])
@login_required
def delete_cable(cable_id):
    row=db.session.get(InfrastructureCable,cable_id)
    if not row:abort(404)
    details={"cable_id":row.cable_id,"type":row.cable_type}
    audit("Infrastructure Cable",row.id,"Delete",old=details,reason=request.form.get("reason"))
    if row.source_device_id:audit("Infrastructure",row.source_device_id,"Cable Removed",old=details,reason=request.form.get("reason"))
    if row.destination_device_id:audit("Infrastructure",row.destination_device_id,"Cable Removed",old=details,reason=request.form.get("reason"))
    db.session.delete(row);db.session.commit();flash("Cable record removed.","success")
    return redirect(url_for("infrastructure.cables"))


@bp.route("/profile/<string:reference>/incidents", methods=["POST"])
@login_required
def add_incident(reference):
    abort(404)
    item=InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:abort(404)
    reported_by=request.form.get("reported_by","").strip()
    problem=request.form.get("problem","").strip()
    if not reported_by or not problem:
        flash("Reported By and Problem are required.","danger")
    else:
        row=InfrastructureIncident(infrastructure=item,reported_by=reported_by,problem=problem,priority=request.form.get("priority") if request.form.get("priority") in INCIDENT_PRIORITIES else "Medium",assigned_engineer=request.form.get("assigned_engineer","").strip() or None,status="Open")
        db.session.add(row);db.session.flush()
        audit("Infrastructure",item.id,"Incident Created",new={"incident":row.incident_number,"priority":row.priority,"problem":problem})
        db.session.commit();flash(f"Incident {row.incident_number} created.","success")
    return redirect(url_for("infrastructure.profile",reference=reference))


@bp.route("/incidents/<int:incident_id>/update", methods=["POST"])
@login_required
def update_incident(incident_id):
    abort(404)
    row=db.session.get(InfrastructureIncident,incident_id)
    if not row:abort(404)
    old={"status":row.status,"root_cause":row.root_cause,"resolution":row.resolution}
    status=request.form.get("status")
    if status not in INCIDENT_STATUSES:abort(400)
    row.status=status;row.assigned_engineer=request.form.get("assigned_engineer","").strip() or row.assigned_engineer
    row.root_cause=request.form.get("root_cause","").strip() or None
    row.resolution=request.form.get("resolution","").strip() or None
    downtime=request.form.get("downtime_minutes",type=int)
    if downtime is not None and downtime < 0:
        flash("Downtime cannot be negative.","danger")
        return redirect(url_for("infrastructure.profile",reference=row.infrastructure.infrastructure_reference))
    row.downtime_minutes=downtime
    row.closed_date=(parse_date(request.form.get("closed_date")) or date.today()) if status in ("Resolved","Closed") else None
    audit("Infrastructure",row.infrastructure_id,"Incident Updated",old=old,new={"incident":row.incident_number,"status":row.status,"root_cause":row.root_cause,"resolution":row.resolution},reason=request.form.get("remarks"))
    db.session.commit();flash(f"Incident {row.incident_number} updated.","success")
    return redirect(url_for("infrastructure.profile",reference=row.infrastructure.infrastructure_reference))


@bp.route("/profile/<string:reference>/verification", methods=["POST"])
@login_required
def add_verification(reference):
    item=InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:abort(404)
    verified_date=parse_date(request.form.get("verified_date"))
    verified_by=request.form.get("verified_by","").strip()
    if not verified_date or not verified_by:
        flash("Verified Date and Verified By are required.","danger")
    else:
        row=InfrastructureVerification(infrastructure=item,verified_date=verified_date,verified_by=verified_by,condition=request.form.get("condition","").strip() or None,next_verification_date=parse_date(request.form.get("next_verification_date")),remarks=request.form.get("remarks","").strip() or None)
        db.session.add(row)
        item.last_verified_date=row.verified_date;item.verified_by=row.verified_by;item.verification_condition=row.condition;item.next_verification_date=row.next_verification_date
        audit("Infrastructure",item.id,"Verification Recorded",new={"verified_date":str(row.verified_date),"verified_by":row.verified_by,"condition":row.condition,"next_verification_date":str(row.next_verification_date) if row.next_verification_date else None},reason=row.remarks)
        db.session.commit();flash("Infrastructure verification recorded.","success")
    return redirect(url_for("infrastructure.profile",reference=reference))


@bp.route("/profile/<string:reference>/qr.png")
@login_required
def qr_image(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    try:
        path, _url, created = ensure_infrastructure_qr(item)
        if created:
            db.session.commit()
    except RuntimeError:
        db.session.rollback()
        abort(503)
    return send_file(path, mimetype="image/png", conditional=True, max_age=3600)


@bp.route("/profile/<string:reference>/qr/download")
@login_required
def qr_download(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    path, _url, created = ensure_infrastructure_qr(item)
    if created:
        db.session.commit()
    return send_file(path, mimetype="image/png", as_attachment=True, download_name=f"{reference}-QR.png")


@bp.route("/profile/<string:reference>/label")
@login_required
def qr_label(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    _path, _url, created = ensure_infrastructure_qr(item)
    if created:
        db.session.commit()
    return render_template("infrastructure_qr_label.html", item=item)


@bp.route("/qr/<string:token>")
@login_required
def qr_resolve(token):
    identity = resolve_infrastructure_token(token)
    if identity:
        return redirect(url_for("infrastructure.profile", reference=identity.infrastructure.infrastructure_reference))
    outside = db.session.execute(
        select(InfrastructureQRIdentity, InfrastructureItem.location_id)
        .join(InfrastructureItem, InfrastructureItem.id == InfrastructureQRIdentity.infrastructure_id)
        .where(InfrastructureQRIdentity.token == token, InfrastructureQRIdentity.active == True)
        .execution_options(lbac_bypass=True)
    ).first()
    if outside and not can_access_location(outside[1]):
        return render_template(
            "access_denied.html", permission="location.view.assigned",
            message="You do not have permission to access this infrastructure record.",
        ), 403
    archived = InfrastructureQRIdentity.query.filter_by(token=token).first()
    if archived:
        return render_template("error.html", title="QR Code Archived", message="This infrastructure QR identity has been replaced. Please use the latest label."), 410
    abort(404)


@bp.route("/profile/<string:reference>/qr/regenerate", methods=["POST"])
@login_required
def qr_regenerate(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    mode = request.form.get("mode", "image")
    reason = request.form.get("reason", "").strip()
    if mode == "identity":
        if normalized_role() != "Super Admin":
            abort(403)
        if request.form.get("identity_confirmation", "").strip() != "REGENERATE IDENTITY":
            flash('Type "REGENERATE IDENTITY" to confirm identity replacement.', "danger")
            return redirect(url_for("infrastructure.profile", reference=reference))
    try:
        regenerate_infrastructure_qr(item, mode, reason, actor=session.get("user_name"))
        db.session.commit()
        flash("Infrastructure QR identity regenerated." if mode == "identity" else "Infrastructure QR image regenerated.", "success")
    except (ValueError, RuntimeError) as exc:
        db.session.rollback()
        flash(str(exc), "danger")
    return redirect(url_for("infrastructure.profile", reference=reference))


@bp.route("/qr-codes", methods=["GET", "POST"])
@login_required
def qr_manage():
    items = InfrastructureItem.query.order_by(InfrastructureItem.infrastructure_reference).all()
    if request.method == "POST":
        action = request.form.get("action")
        reason = request.form.get("reason", "").strip()
        selected = {int(value) for value in request.form.getlist("item_ids") if value.isdigit()}
        targets = items if action in ("missing", "all") else [item for item in items if item.id in selected]
        if action not in ("missing", "selected", "all"):
            flash("Choose a QR action and try again.", "warning")
            return redirect(url_for("infrastructure.qr_manage"))
        if action != "missing" and not reason:
            flash("Regeneration reason is required.", "danger")
            return redirect(url_for("infrastructure.qr_manage"))
        if action == "selected" and not targets:
            flash("Select at least one infrastructure record.", "danger")
            return redirect(url_for("infrastructure.qr_manage"))
        generated = existing = failed = 0
        for item in targets:
            try:
                with db.session.begin_nested():
                    if action == "missing":
                        _path, _url, created = ensure_infrastructure_qr(item, actor=session.get("user_name"), reason="Generate missing infrastructure QR")
                        generated += int(created)
                        existing += int(not created)
                    else:
                        regenerate_infrastructure_qr(item, "image", reason, actor=session.get("user_name"))
                        generated += 1
            except (ValueError, RuntimeError):
                failed += 1
        audit("Infrastructure QR", "Bulk", f"{action.title()} QR Images", new={
            "requested": len(targets), "generated": generated, "existing": existing, "failed": failed,
        }, reason=reason or "Generate missing infrastructure QR")
        db.session.commit()
        flash(f"QR processing complete: {generated} generated, {existing} existing, {failed} failed.", "success" if not failed else "warning")
    ready = {item.id for item in items if qr_metadata(item)["status"] == "Active"}
    return render_template("infrastructure_qr_manage.html", items=items, ready_ids=ready, missing=len(items) - len(ready))


def infrastructure_report_data(kind):
    if kind == "verification":
        return [{
            "Infrastructure ID": row.infrastructure.infrastructure_reference, "Device": row.infrastructure.display_name,
            "Verified Date": row.verified_date, "Verified By": row.verified_by,
            "Condition": row.condition, "Next Verification Date": row.next_verification_date,
            "Remarks": row.remarks,
        } for row in InfrastructureVerification.query.order_by(InfrastructureVerification.verified_date.desc()).all()]
    if kind == "location":
        grouped = {}
        for item in InfrastructureItem.query.order_by(InfrastructureItem.infrastructure_reference).all():
            key = item.location_path
            grouped.setdefault(key, {"Location": key, "Total": 0, "Active": 0, "Under Repair": 0})
            grouped[key]["Total"] += 1
            grouped[key]["Active"] += int(item.status in ("Active", "Installed"))
            grouped[key]["Under Repair"] += int(item.status == "Under Repair")
        return list(grouped.values())
    rows = db.session.query(InfrastructureCategory.name,InfrastructureItem.status,func.count(InfrastructureItem.id)).outerjoin(InfrastructureItem,InfrastructureItem.category_id==InfrastructureCategory.id).group_by(InfrastructureCategory.id,InfrastructureCategory.name,InfrastructureItem.status).order_by(InfrastructureCategory.name,InfrastructureItem.status).all()
    return [{"Category":category,"Status":status or "No Records","Count":count} for category,status,count in rows]


@bp.route("/reports")
@login_required
def reports():
    kind=request.args.get("kind","summary")
    if kind not in ("summary","verification","location"):abort(400)
    rows=infrastructure_report_data(kind)
    titles={"summary":"Infrastructure Summary","verification":"Verification Report","location":"Location Wise Report"}
    if request.args.get("format")=="excel":
        output=io.BytesIO()
        with pd.ExcelWriter(output,engine="openpyxl") as writer:
            pd.DataFrame(rows or [{"Result":"No records"}]).to_excel(writer,index=False,sheet_name="Report",startrow=4)
            sheet=writer.book["Report"];sheet["A1"]="IT Asset Management";sheet["A2"]=titles[kind]
            sheet["A3"]=f"Generated: {datetime.now(ZoneInfo('Asia/Kolkata')):%d %b %Y %I:%M:%S %p}"
            from openpyxl.styles import Font,PatternFill
            sheet["A1"].font=Font(size=18,bold=True,color="203F78");sheet["A2"].font=Font(size=14,bold=True,color="76B82A")
            for cell in sheet[5]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="203F78")
            sheet.freeze_panes="A6";sheet.auto_filter.ref=sheet.dimensions
        output.seek(0)
        return send_file(output,as_attachment=True,download_name=f"infrastructure_{kind}_report.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return render_template("infrastructure_reports.html",kind=kind,title=titles[kind],rows=rows,columns=list(rows[0]) if rows else [],generated_at=datetime.now(ZoneInfo("Asia/Kolkata")))


def export_rows():
    return [{
        "Infrastructure ID": item.infrastructure_reference,
        "Name": item.display_name, "Category": item.category.name,
        "Status": item.status, "Manufacturer": item.manufacturer.name if item.manufacturer else "",
        "Vendor": item.vendor.name if item.vendor else "",
        "Model": item.model, "Serial Number": item.serial_number,
        "Site": item.site, "Floor": item.floor, "Room": item.room,
        "Department": item.department.name if item.department else "",
        "Warranty Expiry": item.warranty_expiry, "AMC Expiry": item.amc_expiry,
        "Management IP": item.effective_management_ip, "Subnet": item.subnet,
        "Gateway": item.gateway, "DNS": item.dns, "VLAN": item.network_vlan,
        "MAC Address": item.mac_address, "DHCP / Static": item.ip_assignment,
        "Last Verified Date": item.last_verified_date, "Verified By": item.verified_by,
        "Verification Condition": item.verification_condition,
        "Next Verification Date": item.next_verification_date,
        "Specifications": "; ".join(f"{key}: {value}" for key, value in item.specifications.items()),
    } for item in filtered_query().order_by(InfrastructureItem.infrastructure_reference).all()]


@bp.route("/export")
@login_required
def export_excel():
    rows = export_rows() or [{"Result": "No infrastructure records"}]
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, index=False, sheet_name="Infrastructure", startrow=4)
        sheet = writer.book["Infrastructure"]
        sheet["A1"] = "IT Asset Management"
        sheet["A2"] = "Infrastructure Register"
        generated_at = datetime.now(ZoneInfo("Asia/Kolkata"))
        sheet["A3"] = f"Generated: {generated_at:%d %b %Y %I:%M:%S %p}"
        from openpyxl.styles import Font, PatternFill
        sheet["A1"].font = Font(size=18, bold=True, color="203F78")
        sheet["A2"].font = Font(size=14, bold=True, color="76B82A")
        for cell in sheet[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203F78")
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = sheet.dimensions
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="infrastructure_register.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/profile/<string:reference>/delete", methods=["POST"])
@login_required
def delete(reference):
    item = InfrastructureItem.query.filter_by(infrastructure_reference=reference).first()
    if not item:
        abort(404)
    reason = request.form.get("reason", "").strip()
    history_count = InfrastructureHistory.query.filter_by(infrastructure_id=item.id).count()
    document_count = InfrastructureDocument.query.filter_by(infrastructure_id=item.id).count()
    enterprise_dependencies = (
        InfrastructureRelationship.query.filter(or_(InfrastructureRelationship.source_id==item.id,InfrastructureRelationship.target_id==item.id)).count()
        + InfrastructurePort.query.filter(or_(InfrastructurePort.infrastructure_id==item.id,InfrastructurePort.connected_device_id==item.id)).count()
        + InfrastructureCable.query.filter(or_(InfrastructureCable.source_device_id==item.id,InfrastructureCable.destination_device_id==item.id)).count()
        + InfrastructureIncident.query.filter_by(infrastructure_id=item.id).count()
        + InfrastructureVerification.query.filter_by(infrastructure_id=item.id).count()
    )
    if not reason:
        flash("Deletion reason is required.", "danger")
    elif history_count > 1 or document_count or enterprise_dependencies:
        flash("Cannot delete infrastructure with operational history, relationships, ports, cables, incidents, verifications, or documents. Use Decommissioning instead.", "danger")
    else:
        path = infrastructure_qr_path(item)
        audit("Infrastructure", item.id, "Delete", old={"reference": reference, "category": item.category.name}, reason=reason)
        db.session.delete(item)
        db.session.commit()
        path.unlink(missing_ok=True)
        flash("Infrastructure record deleted.", "success")
        return redirect(url_for("infrastructure.index"))
    return redirect(url_for("infrastructure.profile", reference=reference))
