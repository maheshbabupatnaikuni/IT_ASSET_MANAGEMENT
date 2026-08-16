from pathlib import Path
import shutil
import sqlite3
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_
from flask import Blueprint, render_template, request, redirect, url_for, flash, current_app, send_from_directory, send_file, abort, session
from werkzeug.utils import secure_filename
from datetime import datetime
import json
import io
from ..extensions import db
from ..models import Asset, Backup, Setting, Department, Location, AssetCategory, Manufacturer, Employee, AuditLog, User, Role, Permission, AssetRequest, LoginHistory, ImportBatch, ImportStagingRow, ImportColumnMapping, ImportValidationError, DynamicField, DynamicFieldValue, ImportSourceLink, FieldProvenance, InfrastructureCategory, InfrastructureSpecificationField, UIFieldConfiguration, utc_now
from ..services.imports import stage_workbook, commit_batch, rollback_batch, dependencies_for_batch, SYSTEM_FIELDS
from ..services.simple_import import import_workbook
from ..services.audit import audit
from ..services.qr_management import ensure_qr_code
from .helpers import login_required, permission_required, normalized_role, can, current_password_valid
from ..services.permissions import PERMISSIONS
from ..services.active_sessions import revoke_active_session
from ..services.ui_configuration import DEFAULT_SETTINGS, FIELD_CATALOG
bp=Blueprint("admin",__name__,url_prefix="/admin")
@bp.before_request
def super_admin_only():
 if not normalized_role():return redirect(url_for("auth.login"))
 if current_app.config.get("SAFE_MODE",True) and request.endpoint in {
  "admin.settings","admin.backup","admin.download_backup","admin.backup_centre",
  "admin.create_recovery_backup","admin.verify_recovery_backup","admin.mark_recovery_tested",
  "admin.prepare_recovery_restore","admin.download_recovery_backup","admin.restore_backup","admin.import_inventory",
  "admin.imports","admin.import_batch","admin.update_import_mapping","admin.update_import_row",
  "admin.commit_import","admin.discard_import","admin.rollback_import","admin.dynamic_fields",
  "admin.import_template","admin.reset_imported_data","admin.delete_user",
 }:
  abort(404)
 endpoint_permission={"admin.users":"user.create" if request.method=="POST" else "user.view","admin.update_user":"user.role.change","admin.reset_user_password":"user.password.reset","admin.toggle_user":"user.toggle","admin.delete_user":"user.delete","admin.login_history":"login_history.view","admin.audit_log":"audit.view","admin.import_inventory":"asset.import","admin.application_configuration":"settings.manage","admin.update_field_configuration":"settings.manage","admin.update_brand_configuration":"settings.manage","admin.create_infrastructure_specification":"settings.manage","admin.update_infrastructure_specification":"settings.manage","admin.backup_centre":"database.backup","admin.create_recovery_backup":"database.backup","admin.verify_recovery_backup":"database.backup","admin.mark_recovery_tested":"database.restore","admin.prepare_recovery_restore":"database.restore","admin.download_recovery_backup":"database.backup"}
 required=endpoint_permission.get(request.endpoint)
 if request.endpoint and request.endpoint.startswith("admin.role") and normalized_role() not in ("Super Admin","Administrator"):abort(403)
 if required and not can(required):abort(403)
 if not required and normalized_role() not in ("Super Admin","Administrator"):abort(403)
 disabled={"admin.imports","admin.import_batch","admin.update_import_mapping","admin.update_import_row","admin.commit_import","admin.discard_import","admin.rollback_import","admin.dynamic_fields","admin.import_template","admin.reset_imported_data"}
 if request.endpoint in {"admin.settings","admin.backup","admin.download_backup","admin.restore_backup"}:abort(404)
 if request.endpoint in disabled:
  flash("The multi-step Excel import workflow is disabled. Use the simple Import Excel page.","warning")
  return redirect(url_for("admin.import_inventory"))
def clean(value):
 return "" if pd.isna(value) else str(value).strip()
def protect_before_change(reason):
 if current_app.config.get("SAFE_MODE",True):
  return {"disabled":True,"reason":reason}
 raise RuntimeError("Recovery tooling is not distributed with this public version.")
def master(model, name):
 name=clean(name) or "Unspecified"
 return model.query.filter_by(name=name).first() or model(name=name)
def import_row(row, index, allow_duplicates=False):
 """Map the supplied HO inventory headers, plus the application's standard headers."""
 source_tag=clean(row.get("asset_tag"))
 tag=source_tag
 if not tag or Asset.query.filter_by(asset_tag=tag).first():
  if not allow_duplicates: return False, f"Row {index}: missing or duplicate Asset Tag"
  tag=f"IMPORTED-{index:04d}"
 source_serial=clean(row.get("serialnumber"))
 serial=source_serial or f"IMP-{tag}"
 if Asset.query.filter_by(serial_number=serial).first(): serial=f"IMP-{index:04d}"
 location=master(Location,row.get("location")); department=master(Department,row.get("department")); category=master(AssetCategory,"Desktop")
 manufacturer=master(Manufacturer,row.get("manufacturer")); db.session.add_all([location,department,category,manufacturer]);db.session.flush()
 username=clean(row.get("username")); employee=None
 if username:
  employee=Employee.query.filter(func.lower(Employee.name)==username.lower()).first()
  if not employee:
   employee=Employee(employee_code=f"IMP-{tag}",name=username,department=department,location=location);db.session.add(employee);db.session.flush()
 a=Asset(asset_tag=tag,source_asset_tag=source_tag or None,category=category,location=location,department=department,employee=employee,username=username or None,
  ip_address=clean(row.get("ip_address")) or None,hostname=clean(row.get("hostname")) or None,mail_domain=clean(row.get("mail_access_domain")) or None,
  manufacturer=manufacturer,model=clean(row.get("model")) or None,serial_number=serial,source_serial_number=source_serial or None,
  operating_system=clean(row.get("operatingsystem")) or None,os_version=clean(row.get("osversion")) or None,processor=clean(row.get("processor")) or None,
  ram=(clean(row.get("ramsizegb")) + " GB") if clean(row.get("ramsizegb")) else None,storage=clean(row.get("hdd/ssd__gb")) or None,
  monitor_make=clean(row.get("monitormake")) or None,monitor_serial_number=clean(row.get("monitorserialno")) or None,monitor_model=clean(row.get("monitormodel")) or None,
  office_mobile=clean(row.get("office_mobile")) or None,mobile_number=clean(row.get("official_number")) or None,
  status="Assigned" if employee else "Available",condition="Good",updated_by="Excel Import")
 db.session.add(a);db.session.flush();ensure_qr_code(a,reason="Automatic QR generation for imported asset");return True, ""
@bp.route("/settings",methods=["GET","POST"])
@login_required
def settings():
 if request.method=="POST":
  for k in ["company_name","theme","default_department","default_location","backup_path"]:
   s=Setting.query.filter_by(key=k).first() or Setting(key=k);s.value=request.form.get(k,"");db.session.add(s)
  db.session.commit();flash("Settings saved.","success")
 values={s.key:s.value for s in Setting.query.all()};return render_template("settings.html",values=values,departments=Department.query.all(),locations=Location.query.all(),backups=Backup.query.order_by(Backup.created_at.desc()).all())
@bp.route("/backup",methods=["POST"])
@login_required
def backup():
 dest=Path(current_app.config["BACKUP_DIR"])/f"itams_{__import__('datetime').datetime.now():%Y%m%d_%H%M%S_%f}.db"
 try:
  source=db.engine.raw_connection(); target=sqlite3.connect(dest); source.connection.backup(target); target.close(); source.close()
  check=sqlite3.connect(dest); integrity=check.execute("PRAGMA integrity_check").fetchone()[0]; check.close()
  verified=integrity=="ok"; db.session.add(Backup(filename=dest.name,size_bytes=dest.stat().st_size,verified=verified,notes="Manual backup"));db.session.commit();flash("Database backup created and verified." if verified else "Backup created but integrity verification failed.","success" if verified else "danger")
 except Exception as exc: db.session.rollback(); dest.unlink(missing_ok=True); flash(f"Backup failed: {exc}","danger")
 return redirect(url_for("admin.settings"))
@bp.route("/backups/<path:name>")
@login_required
def download_backup(name): return send_from_directory(current_app.config["BACKUP_DIR"],name,as_attachment=True)

@bp.route("/backup-centre")
@permission_required("database.backup")
def backup_centre():
 from ..services.recovery import list_backups, recovery_root
 if normalized_role()!="Super Admin":abort(403)
 return render_template("backup_centre.html",backups=list_backups(),recovery_root=recovery_root())

@bp.route("/backup-centre/create",methods=["POST"])
@permission_required("database.backup")
def create_recovery_backup():
 if normalized_role()!="Super Admin":abort(403)
 backup_type=request.form.get("backup_type","").strip()
 reason=request.form.get("reason","").strip()
 if not reason:
  flash("A backup reason is required.","danger")
  return redirect(url_for("admin.backup_centre"))
 try:
  from ..services.recovery import create_backup
  result=create_backup(backup_type,reason)
  audit("Offline Backup",result["backup_id"],"Created",new={"type":backup_type,"version":result["version"],"integrity":result["integrity_result"],"foreign_keys":result["foreign_key_result"],"secondary":result["secondary_status"]},reason=reason)
  db.session.commit()
  flash(f"{backup_type.title()} recovery package created and verified successfully.","success")
 except Exception as exc:
  db.session.rollback();current_app.logger.exception("Recovery package creation failed")
  flash(f"Backup failed safely: {exc}","danger")
 return redirect(url_for("admin.backup_centre"))

@bp.route("/backup-centre/<backup_id>/verify",methods=["POST"])
@permission_required("database.backup")
def verify_recovery_backup(backup_id):
 if normalized_role()!="Super Admin":abort(403)
 try:
  from ..services.recovery import verify_backup
  result=verify_backup(backup_id)
  audit("Offline Backup",backup_id,"Verified",new=result,reason="Manual Backup Centre verification")
  db.session.commit()
  flash("Package, database, foreign keys and archive checksum all passed." if result["valid"] else "Backup verification failed. Do not restore this package.","success" if result["valid"] else "danger")
 except Exception as exc:
  db.session.rollback();current_app.logger.exception("Recovery package verification failed")
  flash(f"Verification failed: {exc}","danger")
 return redirect(url_for("admin.backup_centre"))

@bp.route("/backup-centre/<backup_id>/mark-tested",methods=["POST"])
@permission_required("database.restore")
def mark_recovery_tested(backup_id):
 if normalized_role()!="Super Admin":abort(403)
 notes=request.form.get("test_notes","").strip()
 if not current_password_valid():flash("Your current password is incorrect.","danger")
 elif request.form.get("confirmation","").strip()!="RESTORE TEST PASSED":flash("Type RESTORE TEST PASSED to confirm the completed isolated test.","danger")
 elif not notes:flash("Restore-test notes are required.","danger")
 else:
  try:
   from ..services.recovery import mark_tested
   record=mark_tested(backup_id,session.get("user_name") or "Super Admin",notes)
   audit("Offline Backup",backup_id,"Restore Test Recorded",new={"tested_at":record["tested_at"],"tested_by":record["tested_by"]},reason=notes)
   db.session.commit();flash("Restore-test evidence recorded. Retention will protect this package.","success")
  except Exception as exc:
   db.session.rollback();current_app.logger.exception("Could not record restore test")
   flash(f"Could not record restore test: {exc}","danger")
 return redirect(url_for("admin.backup_centre"))

@bp.route("/backup-centre/<backup_id>/prepare-restore",methods=["POST"])
@permission_required("database.restore")
def prepare_recovery_restore(backup_id):
 if normalized_role()!="Super Admin":abort(403)
 reason=request.form.get("restore_reason","").strip()
 if not current_password_valid():
  flash("Your current password is incorrect.","danger");return redirect(url_for("admin.backup_centre"))
 if request.form.get("confirmation","").strip()!="PREPARE OFFLINE RESTORE":
  flash("Type PREPARE OFFLINE RESTORE to continue.","danger");return redirect(url_for("admin.backup_centre"))
 if not reason:
  flash("A restore reason is required.","danger");return redirect(url_for("admin.backup_centre"))
 try:
  from ..services.recovery import get_backup,verify_backup
  result=verify_backup(backup_id)
  if not result["valid"]:raise ValueError("The selected package did not pass verification.")
  record,_=get_backup(backup_id)
  audit("Offline Backup",backup_id,"Offline Restore Authorized",new={"archive":Path(record["archive"]).name,"type":record["backup_type"],"verified":result},reason=reason)
  db.session.commit()
  return render_template("restore_authorization.html",backup=record,reason=reason)
 except Exception as exc:
  db.session.rollback();current_app.logger.exception("Restore preparation failed")
  flash(f"Restore preparation failed: {exc}","danger");return redirect(url_for("admin.backup_centre"))

@bp.route("/backup-centre/<backup_id>/download")
@permission_required("database.backup")
def download_recovery_backup(backup_id):
 if normalized_role()!="Super Admin":abort(403)
 from ..services.recovery import get_backup
 record,_=get_backup(backup_id);archive=Path(record["archive"])
 if not archive.exists():abort(404)
 audit("Offline Backup",backup_id,"Downloaded",new={"archive":archive.name},reason="Authorised recovery package download")
 db.session.commit()
 return send_file(archive,as_attachment=True,download_name=archive.name)

@bp.route("/audit")
@login_required
def audit_log():
 page=max(request.args.get("page",1,type=int),1); records=AuditLog.query.order_by(AuditLog.timestamp.desc()).paginate(page=page,per_page=100,error_out=False)
 return render_template("audit.html",records=records)

@bp.route("/restore",methods=["POST"])
@login_required
def restore_backup():
 f=request.files.get("backup_file")
 if not f or not f.filename.lower().endswith(".db"):flash("Choose a SQLite .db backup file.","danger");return redirect(url_for("admin.settings"))
 temp=Path(current_app.config["UPLOAD_DIR"])/"restore-validation.db"
 try:
  f.save(temp); candidate=sqlite3.connect(f"file:{temp.as_posix()}?mode=ro",uri=True); result=candidate.execute("PRAGMA integrity_check").fetchone()[0]; tables={r[0] for r in candidate.execute("SELECT name FROM sqlite_master WHERE type='table'")}; candidate.close()
  required={"users","assets","employees","roles"}
  if result!="ok" or not required.issubset(tables):raise ValueError("The selected file is not a valid ITAMS backup")
  database=Path(db.engine.url.database).resolve(); safety=Path(current_app.config["BACKUP_DIR"])/f"pre_restore_{__import__('datetime').datetime.now():%Y%m%d_%H%M%S_%f}.db";shutil.copy2(database,safety)
  db.session.remove();db.engine.dispose();shutil.copy2(temp,database)
  flash("Backup restored successfully. A pre-restore safety copy was retained.","success")
 except Exception as exc:flash(f"Restore failed: {exc}","danger")
 finally:temp.unlink(missing_ok=True)
 return redirect(url_for("admin.settings"))
@bp.route("/import",methods=["GET","POST"])
@login_required
def import_inventory():
 if request.method=="POST":
  f=request.files.get("file")
  if not f or not f.filename.lower().endswith((".xlsx",".xls",".xlsm")):flash("Choose an Excel workbook.","danger");return redirect(url_for("admin.import_inventory"))
  try:
   from ..services.recovery import create_backup
   create_backup("pre-change",f"Before Excel import: {secure_filename(f.filename)}")
   result=import_workbook(f.stream,f.filename)
   return render_template("import_result.html",result=result)
  except Exception as e:db.session.rollback();flash(f"Import failed: {e}","danger")
  return redirect(url_for("admin.import_inventory"))
 return render_template("import_simple.html")

@bp.route("/imports")
@login_required
def imports(): return render_template("imports.html",batches=ImportBatch.query.order_by(ImportBatch.created_at.desc()).all())

@bp.route("/imports/<int:batch_id>")
@login_required
def import_batch(batch_id):
 batch=ImportBatch.query.get_or_404(batch_id);page=max(request.args.get("page",1,type=int),1)
 rows=ImportStagingRow.query.filter_by(batch_id=batch.id).order_by(ImportStagingRow.sheet_id,ImportStagingRow.row_number).paginate(page=page,per_page=100,error_out=False)
 mappings=ImportColumnMapping.query.filter_by(batch_id=batch.id).order_by(ImportColumnMapping.source_column).all()
 return render_template("import_batch.html",batch=batch,rows=rows,mappings=mappings,fields=SYSTEM_FIELDS,dependencies=dependencies_for_batch(batch))

@bp.route("/imports/<int:batch_id>/mapping",methods=["POST"])
@login_required
def update_import_mapping(batch_id):
 batch=ImportBatch.query.get_or_404(batch_id)
 if batch.status in ("Completed","Completed with Warnings","Rolled Back"):flash("Committed mappings cannot be changed.","danger");return redirect(url_for("admin.import_batch",batch_id=batch.id))
 for mapping in ImportColumnMapping.query.filter_by(batch_id=batch.id):
  value=request.form.get(f"mapping_{mapping.id}","dynamic")
  mapping.action="ignore" if value=="ignore" else ("dynamic" if value=="dynamic" else "map");mapping.target_field=value if mapping.action=="map" else None;mapping.reusable=request.form.get("save_mapping")=="on"
 # Reapply mappings to staged source rows, then validate by re-staging the saved workbook.
 path=Path(current_app.config["UPLOAD_DIR"])/batch.stored_filename
 ImportValidationError.query.filter_by(batch_id=batch.id).delete();ImportStagingRow.query.filter_by(batch_id=batch.id).delete()
 for sheet in list(batch.sheets):db.session.delete(sheet)
 batch.valid_rows=batch.invalid_rows=batch.duplicate_rows=batch.total_rows=batch.error_count=0;db.session.commit()
 stage_workbook(batch,path);flash("Mappings saved and workbook revalidated.","success")
 return redirect(url_for("admin.import_batch",batch_id=batch.id))

@bp.route("/imports/rows/<int:row_id>",methods=["POST"])
@login_required
def update_import_row(row_id):
 row=ImportStagingRow.query.get_or_404(row_id);mapped=json.loads(row.mapped_json)
 for key in SYSTEM_FIELDS:
  if key in request.form:mapped[key]=request.form.get(key,"").strip()
 row.mapped_json=json.dumps(mapped);row.proposed_action=request.form.get("proposed_action",row.proposed_action)
 from ..services.imports import validate
 errors,warnings=validate(mapped);row.validation_status="Invalid" if errors else "Valid";row.error_messages="\n".join(errors+warnings);ImportValidationError.query.filter_by(staging_row_id=row.id).delete()
 for message in errors:db.session.add(ImportValidationError(batch_id=row.batch_id,staging_row_id=row.id,severity="error",message=message))
 for message in warnings:db.session.add(ImportValidationError(batch_id=row.batch_id,staging_row_id=row.id,severity="warning",message=message))
 batch=ImportBatch.query.get(row.batch_id);batch.valid_rows=ImportStagingRow.query.filter_by(batch_id=batch.id,validation_status="Valid").count();batch.invalid_rows=ImportStagingRow.query.filter_by(batch_id=batch.id,validation_status="Invalid").count();batch.error_count=batch.invalid_rows;batch.status="Partially Valid" if batch.invalid_rows else "Ready for Review";db.session.commit();flash("Staged row updated and revalidated.","success")
 return redirect(url_for("admin.import_batch",batch_id=row.batch_id))

@bp.route("/imports/<int:batch_id>/commit",methods=["POST"])
@login_required
def commit_import(batch_id):
 batch=ImportBatch.query.get_or_404(batch_id)
 try:commit_batch(batch,session["user_id"]);flash("Import committed atomically to live ITAMS data.","success")
 except Exception as exc:db.session.rollback();batch.status="Rolled Back";db.session.commit();flash(f"Commit rolled back: {exc}","danger")
 return redirect(url_for("admin.import_batch",batch_id=batch.id))

@bp.route("/imports/<int:batch_id>/discard",methods=["POST"])
@login_required
def discard_import(batch_id):
 batch=ImportBatch.query.get_or_404(batch_id)
 if batch.status in ("Completed","Completed with Warnings"):flash("Use rollback for a committed import.","danger")
 else:
  path=Path(current_app.config["UPLOAD_DIR"])/(batch.stored_filename or "");db.session.delete(batch);db.session.commit();path.unlink(missing_ok=True);flash("Staged import discarded. Live data was not changed.","success")
 return redirect(url_for("admin.imports"))

@bp.route("/imports/<int:batch_id>/rollback",methods=["POST"])
@login_required
def rollback_import(batch_id):
 batch=ImportBatch.query.get_or_404(batch_id);ok,deps=rollback_batch(batch)
 flash("Imported source data rolled back." if ok else f"Rollback blocked by operational dependencies: {deps}","success" if ok else "danger")
 return redirect(url_for("admin.import_batch",batch_id=batch.id))

@bp.route("/imports/dynamic-fields")
@login_required
def dynamic_fields():return render_template("dynamic_fields.html",fields=DynamicField.query.order_by(DynamicField.name).all())

@bp.route("/imports/template/<kind>")
@login_required
def import_template(kind):
 if kind not in ("assets","employees","combined"):abort(404)
 wb=Workbook();ws=wb.active;ws.title="Assets" if kind!="employees" else "Employees"
 asset_headers=["Asset Tag","Category","Brand","Model","Serial Number","Hostname","IP Address","Vendor","Purchase Date","Purchase Cost","Warranty Expiry","Invoice Number","Location","Condition","Status","Employee ID","Employee Name","Department","Remarks"]
 employee_headers=["Employee ID","Employee Name","Department","Location","Email"]
 headers=employee_headers if kind=="employees" else asset_headers
 ws.append(headers);ws.append(["ITAM-TAG-0001","Laptop","Dell","Latitude","ITAM-SN-000001","itam-device-01","192.0.2.10","Sample Vendor","2026-04-01",50000,"2029-03-31","ITAM-INV-00001","Corporate Office","Good","Available","ITAM-EMP-000001","Sample Person","Information Technology","Sample row"] if kind!="employees" else ["ITAM-EMP-000001","Sample Person","Information Technology","Corporate Office","person@example.test"])
 for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="312E81")
 ws.freeze_panes="A2"
 if kind=="combined":
  employee=wb.create_sheet("Employees");employee.append(employee_headers);employee.append(["EMP001","Sample Employee","IT","Corporate Office","employee@example.com"])
 instructions=wb.create_sheet("Instructions");instructions.append(["ITAMS Standard Import Template"]);instructions.append(["Required asset field","Asset Tag"]);instructions.append(["Date format","YYYY-MM-DD"]);instructions.append(["Valid asset statuses",", ".join(["Available","Assigned","Blocked","Under Repair","Lost","Scrapped","Disposed"])]);instructions.append(["Workflow","Upload → Mapping → Validation → Review → Super Admin Commit"])
 out=io.BytesIO();wb.save(out);out.seek(0)
 return send_file(out,as_attachment=True,download_name=f"itams_{kind}_import_template.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@bp.route("/imports/reset",methods=["GET","POST"])
@login_required
def reset_imported_data():
 batches=ImportBatch.query.filter(ImportBatch.status.in_(("Completed","Completed with Warnings"))).all();summary={"batches":len(batches),"assets":ImportSourceLink.query.filter(ImportSourceLink.asset_id.isnot(None),ImportSourceLink.detached==False).count(),"employees":ImportSourceLink.query.filter(ImportSourceLink.employee_id.isnot(None),ImportSourceLink.detached==False).count(),"dynamic":DynamicFieldValue.query.count()}
 if request.method=="POST":
  user=User.query.get(session["user_id"]);phrase=request.form.get("confirmation","")
  if not user.check_password(request.form.get("current_password","")) or phrase!="RESET ALL IMPORTED DATA":flash("Password or confirmation phrase is incorrect.","danger")
  else:
   blocked=[] 
   for batch in batches:
    ok,deps=rollback_batch(batch)
    if not ok:blocked.append(f"{batch.filename}: {deps}")
   flash("Imported data reset completed." if not blocked else "Some imports could not be removed because operational dependencies exist: "+"; ".join(blocked),"success" if not blocked else "warning")
  return redirect(url_for("admin.reset_imported_data"))
 return render_template("reset_imports.html",summary=summary)

@bp.route("/users",methods=["GET","POST"])
@login_required
def users():
 if request.method=="POST":
  role=db.session.get(Role,request.form.get("role_id",type=int));username=request.form.get("username","").strip();password=request.form.get("password","")
  location_ids={int(value) for value in request.form.getlist("location_ids") if value.isdigit()}
  locations=Location.query.filter(Location.id.in_(location_ids),Location.active==True).all() if location_ids else []
  if not role or not role.is_active or not username or len(password)<8:flash("Valid username, active role, and an 8-character password are required.","danger")
  elif role.name!="Super Admin" and not locations:flash("Assign at least one active location to this user.","danger")
  elif request.form.get("approval_level")=="Location Admin" and len(locations)!=1:flash("A Location Admin must be assigned to exactly one active location.","danger")
  elif User.query.filter_by(username=username).first():flash("Username already exists.","danger")
  else:
   level="Super Admin" if role.name=="Super Admin" else request.form.get("approval_level","Location User")
   if level not in ("Location User","Location Admin"):level="Location User"
   u=User(username=username,full_name=request.form.get("full_name","").strip() or username,role=role,active=request.form.get("active")=="on",must_change_password=True,approval_level=level,location_scope_enabled=role.name!="Super Admin",locations=locations)
   u.set_password(password);db.session.add(u);db.session.flush();audit("User",u.id,"Created",new={"username":u.username,"role":role.name,"active":u.active,"locations":[x.code for x in locations],"approval_level":level});db.session.commit();flash("User created. They must change the temporary password at first login.","success")
  return redirect(url_for("admin.users"))
 return render_template("users_v2.html",users=User.query.order_by(User.username).all(),roles=Role.query.filter_by(is_active=True).order_by(Role.name).all(),locations=Location.query.filter_by(active=True).order_by(Location.name).all())

def last_active_super(user):
 return user.active and user.role.name in ("IT Admin","Super Admin","Administrator") and User.query.join(Role).filter(User.active==True,Role.name.in_(("IT Admin","Super Admin","Administrator"))).count()<=1

@bp.route("/users/<int:user_id>/update",methods=["POST"])
@login_required
def update_user(user_id):
 user=db.session.get(User,user_id);role=db.session.get(Role,request.form.get("role_id",type=int));username=request.form.get("username",user.username if user else "").strip()
 location_ids={int(value) for value in request.form.getlist("location_ids") if value.isdigit()}
 locations=Location.query.filter(Location.id.in_(location_ids),Location.active==True).all() if location_ids else []
 duplicate=User.query.filter(func.lower(User.username)==username.lower(),User.id!=user_id).first() if username else None
 if not user or not role or not role.is_active:flash("Invalid user or inactive role.","danger")
 elif not username:flash("Username is required.","danger")
 elif duplicate:flash("Username already exists.","danger")
 elif last_active_super(user) and (request.form.get("active")!="on" or role.name!="Super Admin"):flash("The final active Super Admin cannot be deactivated or demoted.","danger")
 elif role.name!="Super Admin" and not locations:flash("Assign at least one active location to this user.","danger")
 elif request.form.get("approval_level")=="Location Admin" and len(locations)!=1:flash("A Location Admin must be assigned to exactly one active location.","danger")
 else:
  old={"username":user.username,"role":user.role.name,"active":user.active,"full_name":user.full_name,"locations":[x.code for x in user.locations],"approval_level":user.approval_level}
  level="Super Admin" if role.name=="Super Admin" else request.form.get("approval_level","Location User")
  if level not in ("Location User","Location Admin"):level="Location User"
  user.username=username;user.full_name=request.form.get("full_name","").strip() or user.full_name;user.role=role;user.active=request.form.get("active")=="on";user.locations=locations;user.approval_level=level;user.location_scope_enabled=role.name!="Super Admin"
  if not user.active:revoke_active_session(user)
  audit("User",user.id,"Updated",old=old,new={"username":user.username,"role":role.name,"active":user.active,"full_name":user.full_name,"locations":[x.code for x in locations],"approval_level":level});db.session.commit();flash("User and location access updated immediately.","success")
 return redirect(url_for("admin.users"))

@bp.route("/locations",methods=["GET","POST"])
@login_required
def locations():
 if request.method=="POST":
  name=request.form.get("name","").strip();code=request.form.get("code","").strip().upper()
  if not name or not code:flash("Location name and code are required.","danger")
  elif Location.query.filter(or_(func.lower(Location.name)==name.lower(),func.lower(Location.code)==code.lower())).first():flash("Location name and code must be unique.","danger")
  else:
   item=Location(name=name,code=code,short_name=request.form.get("short_name","").strip() or None,address=request.form.get("address","").strip() or None,city=request.form.get("city","").strip() or None,state=request.form.get("state","").strip() or None,country=request.form.get("country","").strip() or None,remarks=request.form.get("remarks","").strip() or None,active=True)
   db.session.add(item);db.session.flush();audit("Location",item.id,"Created",new={"name":name,"code":code});db.session.commit();flash("Location created.","success")
  return redirect(url_for("admin.locations"))
 return render_template("locations_admin.html",locations=Location.query.order_by(Location.name).all())

@bp.route("/locations/<int:location_id>/update",methods=["POST"])
@login_required
def update_location(location_id):
 item=db.session.get(Location,location_id)
 if not item:abort(404)
 name=request.form.get("name","").strip();code=request.form.get("code","").strip().upper()
 duplicate=Location.query.filter(Location.id!=item.id,or_(func.lower(Location.name)==name.lower(),func.lower(Location.code)==code.lower())).first()
 if not name or not code:flash("Location name and code are required.","danger")
 elif duplicate:flash("Location name and code must be unique.","danger")
 else:
  old={"name":item.name,"code":item.code,"active":item.active}
  for field in ("short_name","address","city","state","country","remarks"):
   setattr(item,field,request.form.get(field,"").strip() or None)
  item.name=name;item.code=code;item.active=request.form.get("active")=="on"
  audit("Location",item.id,"Updated",old=old,new={"name":name,"code":code,"active":item.active},reason=request.form.get("reason"));db.session.commit();flash("Location updated. Existing linked records were preserved.","success")
 return redirect(url_for("admin.locations"))

@bp.route("/users/<int:user_id>/reset-password",methods=["POST"])
@login_required
def reset_user_password(user_id):
 user=db.session.get(User,user_id);password=request.form.get("password","")
 if not user:flash("User not found.","danger")
 elif len(password)<8:flash("Temporary password must contain at least 8 characters.","danger")
 else:user.set_password(password);user.must_change_password=True;revoke_active_session(user);audit("User",user.id,"Password Reset");db.session.commit();flash("Temporary password saved. The user must change it at next login.","success")
 return redirect(url_for("admin.users"))

@bp.route("/users/<int:user_id>/toggle",methods=["POST"])
@login_required
def toggle_user(user_id):
 user=db.session.get(User,user_id)
 if not user:flash("User not found.","danger")
 elif user.id==session.get("user_id") or last_active_super(user):flash("You cannot deactivate this account.","danger")
 else:
  old=user.active;user.active=not user.active
  if not user.active:revoke_active_session(user)
  audit("User",user.id,"Activated" if user.active else "Deactivated",old={"active":old},new={"active":user.active});db.session.commit();flash("User status updated.","success")
 return redirect(url_for("admin.users"))

@bp.route("/users/<int:user_id>/delete",methods=["POST"])
@login_required
def delete_user(user_id):
 user=db.session.get(User,user_id)
 if not user:flash("User not found.","danger")
 elif user.id==session.get("user_id") or last_active_super(user):flash("You cannot delete this account.","danger")
 else:
  pending=AssetRequest.query.filter(AssetRequest.status.in_(("Pending Approval","Rejected","HR Raised")),or_(AssetRequest.requested_by_id==user.id,AssetRequest.processed_by_id==user.id,AssetRequest.approved_by_id==user.id)).count()
  history=AssetRequest.query.filter(or_(AssetRequest.requested_by_id==user.id,AssetRequest.processed_by_id==user.id,AssetRequest.approved_by_id==user.id)).count()
  if pending:flash("Cannot delete this user. Pending approvals exist. Deactivate the account instead.","danger")
  elif history:flash("Cannot delete this user because assignment history references the account. Deactivate it instead.","danger")
  else:
   old={"username":user.username,"full_name":user.full_name,"role":user.role.name};LoginHistory.query.filter_by(user_id=user.id).delete();db.session.delete(user);audit("User",user.id,"Deleted",old=old);db.session.commit();flash("User deleted.","success")
 return redirect(url_for("admin.users"))

@bp.route("/users/login-history")
@login_required
def login_history():
 return render_template("login_history.html",entries=LoginHistory.query.order_by(LoginHistory.logged_in_at.desc()).limit(500).all())

@bp.route("/roles")
@login_required
def roles():
 return render_template("roles.html",roles=Role.query.order_by(Role.name).all())

@bp.route("/roles/new",methods=["GET","POST"])
@bp.route("/roles/<int:role_id>/edit",methods=["GET","POST"])
@login_required
def role_edit(role_id=None):
 role=db.session.get(Role,role_id) if role_id else None
 if role_id and not role:abort(404)
 if request.method=="POST":
  name=request.form.get("name","").strip();selected=set(request.form.getlist("permissions"))
  if not name:flash("Role name is required.","danger")
  elif Role.query.filter(Role.name==name,Role.id!=(role.id if role else 0)).first():flash("Role name already exists.","danger")
  else:
   creating=role is None
   if creating:role=Role(name=name,created_by=session.get("user_name"),is_system_role=False);db.session.add(role);db.session.flush()
   old={p.code for p in role.permissions}
   if role.is_system_role:name=role.name
   role.name=name;role.description=request.form.get("description","").strip();role.is_active=request.form.get("is_active")=="on"
   if role.name=="Super Admin":role.is_active=True;role.permissions=Permission.query.all()
   else:role.permissions=Permission.query.filter(Permission.code.in_(selected)).all()
   updated={p.code for p in role.permissions}
   audit("Role",role.id,"Created" if creating else "Privileges Updated",old={"permissions":sorted(old)},new={"name":role.name,"active":role.is_active,"permissions":sorted(updated)})
   for code in sorted(updated-old):audit("Role Permission",role.id,"Permission Added",new={"permission":code})
   for code in sorted(old-updated):audit("Role Permission",role.id,"Permission Removed",old={"permission":code})
   if role.users and old!=updated:audit("Role",role.id,"User Access Affected",old={"permissions":sorted(old)},new={"permissions":sorted(updated),"active_users":sum(1 for u in role.users if u.active)})
   db.session.commit();flash("Role and privileges saved.","success");return redirect(url_for("admin.roles"))
 return render_template("role_edit.html",role=role,categories=PERMISSIONS,selected={p.code for p in role.permissions} if role else set())

@bp.route("/roles/<int:role_id>/duplicate",methods=["POST"])
@login_required
def role_duplicate(role_id):
 source=db.session.get(Role,role_id)
 if not source:abort(404)
 name=request.form.get("name","").strip() or f"{source.name} Copy"
 if Role.query.filter_by(name=name).first():flash("Choose a unique duplicate role name.","danger")
 else:
  copy=Role(name=name,description=source.description,is_active=True,is_system_role=False,created_by=session.get("user_name"),permissions=list(source.permissions));db.session.add(copy);db.session.flush();audit("Role",copy.id,"Duplicated",new={"source":source.name,"permissions":[p.code for p in copy.permissions]});db.session.commit();flash("Role duplicated.","success")
 return redirect(url_for("admin.roles"))

@bp.route("/roles/<int:role_id>/toggle",methods=["POST"])
@login_required
def role_toggle(role_id):
 role=db.session.get(Role,role_id)
 if not role:abort(404)
 if role.is_system_role:flash("System roles cannot be deactivated.","danger")
 else:
  old=role.is_active;role.is_active=not role.is_active;audit("Role",role.id,"Activated" if role.is_active else "Deactivated",old={"active":old},new={"active":role.is_active,"affected_active_users":sum(1 for user in role.users if user.active)});db.session.commit()
 return redirect(url_for("admin.roles"))


@bp.route("/application-configuration")
@login_required
def application_configuration():
 if normalized_role()!="Super Admin":abort(403)
 selected=request.args.get("module","assets")
 if selected not in FIELD_CATALOG:selected="assets"
 fields=UIFieldConfiguration.query.filter_by(module=selected).order_by(UIFieldConfiguration.sort_order,UIFieldConfiguration.id).all()
 settings={row.key:row.value for row in Setting.query.filter(Setting.key.in_(DEFAULT_SETTINGS)).all()}
 infrastructure_categories=InfrastructureCategory.query.filter_by(active=True).order_by(InfrastructureCategory.name).all() if selected=="infrastructure" else []
 selected_category_id=request.args.get("category",type=int) or (infrastructure_categories[0].id if infrastructure_categories else None)
 infrastructure_fields=InfrastructureSpecificationField.query.filter_by(category_id=selected_category_id).order_by(InfrastructureSpecificationField.sort_order,InfrastructureSpecificationField.id).all() if selected_category_id else []
 return render_template("application_configuration.html",modules=FIELD_CATALOG,selected_module=selected,fields=fields,settings=settings,defaults=DEFAULT_SETTINGS,infrastructure_categories=infrastructure_categories,selected_category_id=selected_category_id,infrastructure_fields=infrastructure_fields)


@bp.route("/application-configuration/fields/<int:field_id>",methods=["POST"])
@login_required
def update_field_configuration(field_id):
 if normalized_role()!="Super Admin":abort(403)
 row=db.session.get(UIFieldConfiguration,field_id)
 if not row:abort(404)
 old={"label":row.label,"visible":row.visible,"required":row.required,"sort_order":row.sort_order,"placeholder":row.placeholder,"help_text":row.help_text}
 if request.form.get("action")=="reset":
  protect_before_change(f"Before UI field configuration reset: {row.module}.{row.field_code}")
  row.label=row.default_label;row.visible=True;row.required=row.protected;row.sort_order=request.form.get("default_order",row.sort_order,type=int);row.placeholder=None;row.help_text=None
  action="Field Configuration Reset"
 else:
  label=request.form.get("label","").strip()
  if not label:
   flash("Field label cannot be blank.","danger");return redirect(url_for("admin.application_configuration",module=row.module))
  protect_before_change(f"Before UI field configuration change: {row.module}.{row.field_code}")
  row.label=label
  row.visible=True if row.protected else request.form.get("visible")=="on"
  row.required=True if row.protected else request.form.get("required")=="on"
  row.sort_order=max(0,min(request.form.get("sort_order",100,type=int),9999))
  row.placeholder=request.form.get("placeholder","").strip() or None
  row.help_text=request.form.get("help_text","").strip() or None
  action="Field Configuration Updated"
 row.updated_by=session.get("user_name")
 db.session.flush();audit("UI Field Configuration",row.id,action,old=old,new={"module":row.module,"field_code":row.field_code,"label":row.label,"visible":row.visible,"required":row.required,"sort_order":row.sort_order,"placeholder":row.placeholder,"help_text":row.help_text});db.session.commit()
 flash(f"{row.label} configuration saved.","success")
 return redirect(url_for("admin.application_configuration",module=row.module))


@bp.route("/application-configuration/branding",methods=["POST"])
@login_required
def update_brand_configuration():
 if normalized_role()!="Super Admin":abort(403)
 old={key:(Setting.query.filter_by(key=key).first().value if Setting.query.filter_by(key=key).first() else None) for key in DEFAULT_SETTINGS}
 values={}
 for key,default in DEFAULT_SETTINGS.items():
  value=request.form.get(key,"").strip() or default
  if key.endswith("_color") and not re.fullmatch(r"#[0-9a-fA-F]{6}",value):
   flash("Brand colours must use the format #RRGGBB.","danger");return redirect(url_for("admin.application_configuration"))
  values[key]=value
 protect_before_change("Before application branding configuration change")
 values={}
 for key,default in DEFAULT_SETTINGS.items():
  value=request.form.get(key,"").strip() or default
  row=Setting.query.filter_by(key=key).first() or Setting(key=key);row.value=value;db.session.add(row);values[key]=value
 audit("Application Configuration","Branding","Branding Updated",old=old,new=values);db.session.commit();flash("Application branding and login text updated.","success")
 return redirect(url_for("admin.application_configuration"))


@bp.route("/application-configuration/infrastructure-specifications",methods=["POST"])
@login_required
def create_infrastructure_specification():
 if normalized_role()!="Super Admin":abort(403)
 category=db.session.get(InfrastructureCategory,request.form.get("category_id",type=int))
 label=request.form.get("label","").strip()
 code=re.sub(r"[^a-z0-9]+","_",(request.form.get("field_code") or label).strip().lower()).strip("_")
 if not category or not label or not code:
  flash("Choose a category and enter a field name.","danger")
 elif InfrastructureSpecificationField.query.filter_by(category_id=category.id,field_code=code).first():
  flash("That field code already exists for this category. Edit or reactivate the existing field.","danger")
 else:
  protect_before_change(f"Before infrastructure specification field addition: {category.name}.{code}")
  row=InfrastructureSpecificationField(category=category,field_code=code,label=label,active=True,sort_order=request.form.get("sort_order",100,type=int),created_by=session.get("user_name"),updated_by=session.get("user_name"));db.session.add(row);db.session.flush();audit("Infrastructure Specification",row.id,"Field Added",new={"category":category.name,"field_code":code,"label":label,"sort_order":row.sort_order});db.session.commit();flash(f"{label} added to {category.name}.","success")
 return redirect(url_for("admin.application_configuration",module="infrastructure",category=category.id if category else None))


@bp.route("/application-configuration/infrastructure-specifications/<int:field_id>",methods=["POST"])
@login_required
def update_infrastructure_specification(field_id):
 if normalized_role()!="Super Admin":abort(403)
 row=db.session.get(InfrastructureSpecificationField,field_id)
 if not row:abort(404)
 label=request.form.get("label","").strip()
 if not label:
  flash("Field label cannot be blank.","danger")
 else:
  protect_before_change(f"Before infrastructure specification field update: {row.category.name}.{row.field_code}")
  old={"label":row.label,"active":row.active,"sort_order":row.sort_order};row.label=label;row.active=request.form.get("active")=="on";row.sort_order=max(0,min(request.form.get("sort_order",100,type=int),9999));row.updated_by=session.get("user_name");audit("Infrastructure Specification",row.id,"Field Updated",old=old,new={"category":row.category.name,"field_code":row.field_code,"label":row.label,"active":row.active,"sort_order":row.sort_order},reason=request.form.get("reason") or "Frontend configuration change");db.session.commit();flash(f"{row.label} configuration saved. Existing stored values were retained.","success")
 return redirect(url_for("admin.application_configuration",module="infrastructure",category=row.category_id))
